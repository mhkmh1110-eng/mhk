#!/usr/bin/env python3
"""인터넷등기소 부동산 고유번호 취합 프로그램 v1.5

폴더에 주소 목록이 담긴 엑셀 파일을 두고 실행하면:

  1) 크롬 자동 시작 → 인터넷등기소(https://www.iros.go.kr) 접속 (로그인 없이 진행)
  2) 메인화면의 "부동산" 통합 검색창에 주소 질의어를 자동 입력하여 검색
     - 주소 질의어 = ["번지" 앞에 있는 부분] + " " + "마지막 N동 N호"
     - 예: "경기도 고양시 일산동구 마두동 732번지 백마한양아파트 310동 1005호"
       → "경기도 고양시 일산동구 마두동 732 310동 1005호"
  3) 검색 결과 페이지(간편검색 탭)에서 "전체 N 건" 카운트를 읽고:
     - 정확히 1건이면 결과 테이블에서
       (부동산고유번호 / 부동산구분 / 부동산표시) 3개 컬럼 값을 추출하여
       엑셀의 해당 작업번호 행 우측에 3개 셀로 기록.
     - 0건 또는 2건 이상이면 비고만 남기고 다음 작업으로 진행.
  4) 한 건 끝나면 인터넷등기소 메인으로 복귀하여 다음 주소 검색 진행.

기존 등기부등본 발급 프로그램(add_to_payment_v5_8) 의 검증된 인프라를
재사용:
  - 크롬 옵션, 경량 프로필 복사 (보안 모듈 인식)
  - ChromeDriver 다단계 fallback (SSL 인터셉션 회피)
  - processbar 오버레이가 클릭을 가로채는 경우 제거 후 재클릭
  - 빈 화면 자동 새로고침
  - 보안 프로그램 설치 페이지 자동 복구
  - Windows '응용 프로그램 오류' 팝업 자동 닫기 + 30초 휴식 워치독
  - 절전/화면보호기 방지
  - 진행상황 오버레이 + 일시정지/중단 버튼 (텍스트 자동 줄바꿈)

엑셀 형식 (자동 감지):
  - 헤더 행(1행)에 "주소" 컬럼이 있어야 함.
  - 작업번호 컬럼이 있으면 로그/출력 위치 기준으로 사용 (없어도 동작).
  - 결과는 "부동산고유번호 / 부동산구분 / 부동산표시 / 비고" 헤더가 있으면
    그 컬럼에 기록. 없으면 주소 컬럼 우측에 4개 컬럼을 자동 생성.

사용법:
  - Windows 에서 collect_property_ids_v1_5.py 더블클릭 (가장 간단)
  - python collect_property_ids_v1_5.py
  - python collect_property_ids_v1_5.py --file 주소목록.xlsx
  - python collect_property_ids_v1_5.py --start-from 5
"""

import argparse
import ctypes
from copy import copy
import os
import platform
import queue
import re
import shutil
import sys
import threading
import time
import logging
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.action_chains import ActionChains
from selenium.common.exceptions import (
    TimeoutException,
    NoSuchElementException,
    NoAlertPresentException,
    UnexpectedAlertPresentException,
    ElementClickInterceptedException,
    ElementNotInteractableException,
    StaleElementReferenceException,
)
from webdriver_manager.chrome import ChromeDriverManager

# 기업/기관망의 SSL 인터셉션으로 인해 chromedriver 버전 조회가
# 자체서명 인증서 오류로 실패하는 문제 회피.
os.environ.setdefault("WDM_SSL_VERIFY", "0")


# ─── 설정 ──────────────────────────────────────────────────
IROS_BASE_URL = "https://www.iros.go.kr"
IROS_INDEX_URL = "https://www.iros.go.kr/index.jsp"
ELEMENT_WAIT_TIMEOUT = 15
DEFAULT_FILE = "주소목록.xlsx"

# 한 주소를 처리할 때 검색~결과 추출까지 최대 시도 라운드 수.
# v1.3: 서버 오류(검색창 미감지/결과 페이지 미진입/테이블 추출 실패)만 재시도.
#       결과 없음(0건) 또는 다건(2건+)은 재시도 없이 바로 다음 작업으로 진행.
MAX_RETRY_PER_ITEM = 10
RETRY_DELAY_SEC = 2.0

# 검색 후 결과 페이지(간편검색 탭 + "부동산 소재지번 검색 결과" 영역)가
# 렌더될 때까지의 최대 폴링 대기 시간 (초). 등기소 서버 지연 흡수.
SEARCH_RESULT_WAIT_TIMEOUT = 20
SEARCH_RESULT_POLL = 0.3

# 결과 추출 시 결과 표가 안정화될 때까지의 추가 안정화 폴링 (초).
RESULT_STABILIZE_WAIT = 1.0

# ─── 엑셀 컬럼 헤더 후보 (자동 감지용, 대소문자/공백 무시 매칭) ──────
HEADER_ADDRESS_KEYWORDS    = ("주소",)
HEADER_TASK_NO_KEYWORDS    = ("작업번호", "순번", "번호", "no", "no.", "#")
HEADER_UNIQUE_NO_KEYWORDS  = ("부동산고유번호", "고유번호")
HEADER_KIND_KEYWORDS       = ("부동산구분", "구분")
HEADER_REPR_KEYWORDS       = ("부동산표시", "표시")
HEADER_NOTE_KEYWORDS       = ("비고", "결과", "메모")

# 결과 출력 헤더가 없으면 새로 만들 때 사용할 라벨
DEFAULT_HEADER_UNIQUE = "부동산고유번호"
DEFAULT_HEADER_KIND   = "부동산구분"
DEFAULT_HEADER_REPR   = "부동산표시"
DEFAULT_HEADER_NOTE   = "비고"

# ─── 로깅 설정 ─────────────────────────────────────────────
log_dir = "logs"
os.makedirs(log_dir, exist_ok=True)
log_filename = os.path.join(
    log_dir, f"collect_property_ids_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"
)
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.FileHandler(log_filename, encoding="utf-8"),
        logging.StreamHandler(sys.stdout),
    ],
)
logger = logging.getLogger(__name__)


# ═══════════════════════════════════════════════════════════
# 진행상황 오버레이 (add_to_payment_v5_8 에서 가져온 동일 컴포넌트)
# ═══════════════════════════════════════════════════════════
class ProgressOverlay:
    """우측 하단 플로팅 로그 오버레이."""

    def __init__(self, max_lines=10, width=540, height=270, margin=24, alpha=0.7,
                 controller=None):
        self.max_lines = max_lines
        self.width = width
        self.height = height
        self.margin = margin
        self.alpha = alpha
        self.controller = controller
        self._q = queue.Queue()
        self._thread = None
        self._ready = threading.Event()
        self._stopped = threading.Event()

    @classmethod
    def start(cls, **kwargs):
        obj = cls(**kwargs)
        obj._thread = threading.Thread(target=obj._run_tk, name="overlay-tk", daemon=True)
        obj._thread.start()
        obj._ready.wait(timeout=2.0)
        return obj

    def log(self, level, message):
        self._q.put(("log", (level, message)))

    def set_status(self, text):
        self._q.put(("status", text))

    def set_progress(self, current, total):
        self._q.put(("progress", (current, total)))

    def set_pause_indicator(self, paused):
        self._q.put(("pause_state", bool(paused)))

    def set_stopped(self):
        self._q.put(("stopped", None))

    def attach_to_logger(self, target_logger, level=logging.INFO,
                         fmt="%(asctime)s %(message)s", datefmt="%H:%M:%S"):
        h = TkOverlayHandler(self)
        h.setLevel(level)
        h.setFormatter(logging.Formatter(fmt, datefmt=datefmt))
        target_logger.addHandler(h)
        return h

    def stop(self):
        if self._stopped.is_set():
            return
        self._stopped.set()
        self._q.put(("__quit__", None))
        if self._thread and self._thread.is_alive():
            self._thread.join(timeout=1.5)

    def _run_tk(self):
        try:
            import tkinter as tk
            from tkinter import ttk
        except Exception as e:
            print(f"[overlay] Tkinter 사용 불가, 오버레이 비활성화: {e}")
            self._ready.set()
            return

        root = tk.Tk()
        root.title("진행상황")
        root.overrideredirect(True)
        root.attributes("-topmost", True)
        try:
            root.attributes("-alpha", self.alpha)
        except tk.TclError:
            pass

        root.update_idletasks()
        screen_w = root.winfo_screenwidth()
        screen_h = root.winfo_screenheight()
        x = screen_w - self.width - self.margin
        y = screen_h - self.height - self.margin - 48
        if y < 0:
            y = self.margin
        root.geometry(f"{self.width}x{self.height}+{x}+{y}")

        BG        = "#1e1e1e"
        BG_HEADER = "#2d2d30"
        FG        = "#e8e8e8"
        FG_MUTED  = "#9a9a9a"
        ACCENT    = "#4ea1ff"
        ERR       = "#ff6b6b"
        WARN      = "#ffb86b"

        root.configure(bg=BG)

        title_bar = tk.Frame(root, bg=BG_HEADER, height=26)
        title_bar.pack(fill="x", side="top")
        title_bar.pack_propagate(False)
        title_lbl = tk.Label(title_bar, text="● 진행상황",
                             bg=BG_HEADER, fg=FG,
                             font=("Segoe UI", 9, "bold"), padx=10)
        title_lbl.pack(side="left")
        close_btn = tk.Label(title_bar, text="✕",
                             bg=BG_HEADER, fg=FG_MUTED,
                             font=("Segoe UI", 10, "bold"),
                             padx=10, cursor="hand2")
        close_btn.pack(side="right")

        def _close(_e=None):
            try:
                root.destroy()
            except Exception:
                pass
        close_btn.bind("<Button-1>", _close)
        close_btn.bind("<Enter>", lambda e: close_btn.config(fg=ERR))
        close_btn.bind("<Leave>", lambda e: close_btn.config(fg=FG_MUTED))

        drag_state = {"x": 0, "y": 0}
        def _start_move(e):
            drag_state["x"], drag_state["y"] = e.x, e.y
        def _do_move(e):
            nx = root.winfo_x() + (e.x - drag_state["x"])
            ny = root.winfo_y() + (e.y - drag_state["y"])
            root.geometry(f"+{nx}+{ny}")
        for w in (title_bar, title_lbl):
            w.bind("<Button-1>", _start_move)
            w.bind("<B1-Motion>", _do_move)

        status_frame = tk.Frame(root, bg=BG)
        status_frame.pack(fill="x", side="top", padx=10, pady=(6, 0))
        status_lbl = tk.Label(status_frame, text="대기 중…",
                              bg=BG, fg=ACCENT,
                              font=("Segoe UI", 9), anchor="w",
                              justify="left", wraplength=self.width - 24)
        status_lbl.pack(fill="x", side="top")
        progress_lbl = tk.Label(status_frame, text="",
                                bg=BG, fg=FG_MUTED,
                                font=("Segoe UI", 8), anchor="w")
        progress_lbl.pack(fill="x", side="top")

        style = ttk.Style()
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure(
            "overlay.Horizontal.TProgressbar",
            troughcolor=BG_HEADER, background=ACCENT,
            bordercolor=BG_HEADER, lightcolor=ACCENT, darkcolor=ACCENT,
            thickness=6,
        )
        pbar = ttk.Progressbar(status_frame,
                               style="overlay.Horizontal.TProgressbar",
                               mode="determinate", maximum=100)
        pbar.pack(fill="x", side="top", pady=(4, 4))

        BTN_BG        = "#2d2d30"
        BTN_BG_HOVER  = "#3e3e42"
        BTN_BG_ACTIVE = "#094771"
        BTN_FG        = "#e8e8e8"
        BTN_FG_DIS    = "#5a5a5a"

        btn_frame = tk.Frame(root, bg=BG)
        btn_frame.pack(fill="x", side="top", padx=10, pady=(2, 4))

        ctrl = self.controller

        def _make_btn(parent, text, on_click):
            b = tk.Label(parent, text=text, bg=BTN_BG, fg=BTN_FG,
                         font=("Segoe UI", 9), padx=10, pady=3,
                         cursor="hand2")
            b._enabled = True
            b._base_bg = BTN_BG
            def _enter(_e):
                if b._enabled:
                    b.config(bg=BTN_BG_HOVER)
            def _leave(_e):
                if b._enabled:
                    b.config(bg=b._base_bg)
            def _click(_e):
                if b._enabled:
                    try:
                        on_click()
                    except Exception as ex:
                        logger.debug(f"button click error: {ex}")
            b.bind("<Enter>", _enter)
            b.bind("<Leave>", _leave)
            b.bind("<Button-1>", _click)
            return b

        def _set_enabled(btn, enabled, base_bg=None):
            btn._enabled = enabled
            if base_bg is not None:
                btn._base_bg = base_bg
            if enabled:
                btn.config(fg=BTN_FG, bg=btn._base_bg, cursor="hand2")
            else:
                btn.config(fg=BTN_FG_DIS, bg=BTN_BG, cursor="arrow")

        def _on_pause():
            if ctrl is not None:
                ctrl.pause()
                logger.info("[UI] 일시정지 요청 — 현재 작업이 안전 지점에 도달하면 멈춥니다")

        def _on_resume():
            if ctrl is not None:
                ctrl.resume()
                logger.info("[UI] 재개 요청")

        def _on_stop():
            if ctrl is not None:
                ctrl.request_stop()
                logger.warning("[UI] 중단 요청 — 현재 건을 마치고 안전 종료합니다")

        pause_btn  = _make_btn(btn_frame, "⏸ 일시정지", _on_pause)
        resume_btn = _make_btn(btn_frame, "▶ 재개",     _on_resume)
        stop_btn   = _make_btn(btn_frame, "⏹ 중단",     _on_stop)
        pause_btn.pack(side="left", padx=(0, 4))
        resume_btn.pack(side="left", padx=(0, 4))
        stop_btn.pack(side="left", padx=(0, 4))

        if ctrl is None:
            _set_enabled(pause_btn, False)
            _set_enabled(resume_btn, False)
            _set_enabled(stop_btn, False)
        else:
            _set_enabled(pause_btn, True, BTN_BG)
            _set_enabled(resume_btn, False)
            _set_enabled(stop_btn, True, "#5a2020")

        log_frame = tk.Frame(root, bg=BG)
        log_frame.pack(fill="both", expand=True, padx=10, pady=(2, 10))
        txt = tk.Text(log_frame, bg=BG, fg=FG,
                      font=("Consolas", 9),
                      borderwidth=0, highlightthickness=0,
                      wrap="word", state="disabled", cursor="arrow")
        txt.pack(fill="both", expand=True)
        txt.tag_configure("INFO",    foreground=FG)
        txt.tag_configure("DEBUG",   foreground=FG_MUTED)
        txt.tag_configure("WARNING", foreground=WARN)
        txt.tag_configure("ERROR",   foreground=ERR)

        self._ready.set()

        def _pump():
            drained = 0
            while drained < 30:
                try:
                    kind, payload = self._q.get_nowait()
                except queue.Empty:
                    break
                drained += 1
                try:
                    if kind == "__quit__":
                        try:
                            root.destroy()
                        except Exception:
                            pass
                        return
                    elif kind == "log":
                        level, message = payload
                        tag = level if level in ("INFO", "DEBUG", "WARNING", "ERROR") else "INFO"
                        txt.configure(state="normal")
                        txt.insert("end", message.rstrip() + "\n", tag)
                        line_count = int(txt.index("end-1c").split(".")[0])
                        if line_count > self.max_lines:
                            txt.delete("1.0", f"{line_count - self.max_lines + 1}.0")
                        txt.see("end")
                        txt.configure(state="disabled")
                    elif kind == "status":
                        status_lbl.config(text=str(payload))
                    elif kind == "progress":
                        cur, tot = payload
                        if tot > 0:
                            pct = max(0, min(100, int(cur * 100 / tot)))
                            pbar["value"] = pct
                            progress_lbl.config(text=f"[{cur}/{tot}] ({pct}%)")
                        else:
                            pbar["value"] = 0
                            progress_lbl.config(text="")
                    elif kind == "pause_state":
                        paused = bool(payload)
                        if paused:
                            _set_enabled(pause_btn, False)
                            _set_enabled(resume_btn, True, BTN_BG_ACTIVE)
                        else:
                            _set_enabled(pause_btn, True, BTN_BG)
                            _set_enabled(resume_btn, False)
                    elif kind == "stopped":
                        _set_enabled(pause_btn, False)
                        _set_enabled(resume_btn, False)
                        _set_enabled(stop_btn, False)
                except Exception:
                    pass
            try:
                root.after(80, _pump)
            except Exception:
                pass

        root.after(50, _pump)
        try:
            root.mainloop()
        except Exception:
            pass


class TkOverlayHandler(logging.Handler):
    """logging → ProgressOverlay 브릿지 핸들러."""

    def __init__(self, overlay):
        super().__init__()
        self._overlay = overlay

    def emit(self, record):
        try:
            msg = self.format(record)
            self._overlay.log(record.levelname, msg)
        except Exception:
            self.handleError(record)


_overlay = None


# ═══════════════════════════════════════════════════════════
# 일시정지/중단 컨트롤러
# ═══════════════════════════════════════════════════════════
class StopRequested(Exception):
    """사용자가 오버레이에서 '중단' 을 눌렀을 때 발생."""
    pass


class PauseController:
    def __init__(self):
        self._paused = threading.Event()
        self._stopping = threading.Event()
        self._on_state_change = None

    def set_state_callback(self, fn):
        self._on_state_change = fn

    @property
    def is_paused(self):
        return self._paused.is_set()

    @property
    def is_stopping(self):
        return self._stopping.is_set()

    def pause(self):
        if self._stopping.is_set():
            return
        self._paused.set()
        if self._on_state_change:
            try:
                self._on_state_change(True)
            except Exception:
                pass

    def resume(self):
        self._paused.clear()
        if self._on_state_change:
            try:
                self._on_state_change(False)
            except Exception:
                pass

    def request_stop(self):
        self._stopping.set()
        self._paused.clear()

    def checkpoint(self):
        if self._stopping.is_set():
            raise StopRequested()
        if self._paused.is_set():
            logger.info("⏸ 일시정지 (오버레이 '재개' 버튼을 누르면 계속)")
            try:
                overlay_status("⏸ 일시정지됨")
            except Exception:
                pass
            while self._paused.is_set() and not self._stopping.is_set():
                time.sleep(0.2)
            if self._stopping.is_set():
                raise StopRequested()
            logger.info("▶ 재개")

    def sleep(self, seconds):
        end = time.time() + seconds
        while True:
            if self._stopping.is_set():
                raise StopRequested()
            if self._paused.is_set():
                self.checkpoint()
                end = time.time() + seconds
                continue
            remaining = end - time.time()
            if remaining <= 0:
                return
            time.sleep(min(0.2, remaining))


_pause = None
_popup_watcher = None


# ═══════════════════════════════════════════════════════════
# Windows '응용 프로그램 오류' 팝업 자동 닫기 워치독
#   - 1회 발생 시 현재 건 마무리 후 30초 휴식 (보안 모듈 회복 시간 확보)
# ═══════════════════════════════════════════════════════════
class PopupWatcher:
    _TITLE_KEYWORDS = (
        "응용 프로그램 오류",
        "Application Error",
        "응용 프로그램이 응답하지 않습니다",
    )
    _BUTTON_TEXTS = ("확인", "OK", "예", "Yes", "닫기", "Close")

    _BM_CLICK   = 0x00F5
    _WM_CLOSE   = 0x0010
    _WM_GETTEXT = 0x000D
    _WM_GETTEXTLENGTH = 0x000E
    _GW_HWNDNEXT = 2
    _GW_CHILD    = 5

    def __init__(self, poll_interval=5.0, rest_seconds=30, rest_threshold=1):
        self.poll_interval = float(poll_interval)
        self.rest_seconds = int(rest_seconds)
        self.rest_threshold = int(rest_threshold)

        self._stop_evt = threading.Event()
        self._thread = None
        self._enabled = (platform.system() == "Windows")

        self._lock = threading.Lock()
        self.dismiss_count = 0
        self._pending_rest = 0
        self.rest_granted_total = 0
        self.rest_consumed_total = 0

        self._recent_handled = {}

    def start(self):
        if not self._enabled:
            logger.info("팝업 워치독 건너뜀 (Windows 아님)")
            return False
        if self._thread is not None and self._thread.is_alive():
            return True
        self._stop_evt.clear()
        self._thread = threading.Thread(
            target=self._run, name="popup-watcher", daemon=True
        )
        self._thread.start()
        logger.info(
            f"팝업 워치독 시작 (폴링 {self.poll_interval:.0f}초, "
            f"발생 시 현재 건 마무리 후 {self.rest_seconds}초 휴식)"
        )
        return True

    def stop(self):
        if not self._enabled:
            return
        self._stop_evt.set()
        if self._thread is not None:
            self._thread.join(timeout=2)
        self._thread = None

    def consume_rest_if_due(self):
        if not self._enabled:
            return False
        with self._lock:
            if self._pending_rest <= 0:
                return False
            self._pending_rest -= 1
            self.rest_consumed_total += 1
            secs = self.rest_seconds
        logger.warning(
            f"[팝업워치독] OS 팝업 후속 휴식 {secs}초 — 보안 모듈 회복 시간 확보"
        )
        try:
            if _overlay is not None:
                _overlay.set_status(f"⏸ 휴식 {secs}초")
        except Exception:
            pass
        if _pause is not None:
            _pause.sleep(secs)
        else:
            time.sleep(secs)
        return True

    def _run(self):
        try:
            user32 = ctypes.windll.user32
        except Exception as e:
            logger.debug(f"팝업 워치독 비활성 (user32 로드 실패: {e})")
            return

        while not self._stop_evt.is_set():
            try:
                self._scan_once(user32)
            except Exception as e:
                logger.debug(f"팝업 워치독 스캔 예외 (무시): {e}")
            end = time.time() + self.poll_interval
            while time.time() < end:
                if self._stop_evt.is_set():
                    return
                time.sleep(0.2)

    def _scan_once(self, user32):
        now = time.time()
        if self._recent_handled:
            self._recent_handled = {
                h: t for h, t in self._recent_handled.items() if now - t < 60
            }

        EnumWindowsProc = ctypes.WINFUNCTYPE(
            ctypes.c_bool, ctypes.c_void_p, ctypes.c_void_p
        )
        hits = []

        def _cb(hwnd, _lparam):
            try:
                if not user32.IsWindowVisible(hwnd):
                    return True
                cls = ctypes.create_unicode_buffer(64)
                user32.GetClassNameW(hwnd, cls, 64)
                if cls.value != "#32770":
                    return True
                title = self._get_window_text(user32, hwnd)
                if not title:
                    return True
                low = title.lower()
                if not any(k.lower() in low for k in self._TITLE_KEYWORDS):
                    return True
                hits.append((hwnd, title))
            except Exception:
                pass
            return True

        user32.EnumWindows(EnumWindowsProc(_cb), 0)

        for hwnd, title in hits:
            if hwnd in self._recent_handled:
                continue
            self._dismiss(user32, hwnd, title)
            self._recent_handled[hwnd] = time.time()

    def _get_window_text(self, user32, hwnd):
        try:
            length = user32.SendMessageW(hwnd, self._WM_GETTEXTLENGTH, 0, 0)
            if length <= 0:
                return ""
            buf = ctypes.create_unicode_buffer(length + 2)
            user32.SendMessageW(hwnd, self._WM_GETTEXT, length + 1, ctypes.byref(buf))
            return buf.value or ""
        except Exception:
            return ""

    def _collect_body_text(self, user32, hwnd):
        texts = []
        try:
            child = user32.GetWindow(hwnd, self._GW_CHILD)
            seen = 0
            while child and seen < 30:
                seen += 1
                t = self._get_window_text(user32, child)
                if t:
                    texts.append(t)
                child = user32.GetWindow(child, self._GW_HWNDNEXT)
        except Exception:
            pass
        return " | ".join(texts)

    def _dismiss(self, user32, hwnd, title):
        body = self._collect_body_text(user32, hwnd)
        logger.warning(
            f"[팝업워치독] OS 오류 팝업 감지 → 자동 닫기 시도\n"
            f"  · 제목: {title}\n"
            f"  · 내용: {body[:300]}"
        )

        clicked = self._click_button(user32, hwnd)
        if not clicked:
            try:
                user32.PostMessageW(hwnd, self._WM_CLOSE, 0, 0)
                logger.info("[팝업워치독] 버튼 미검출 → WM_CLOSE 폴백 전송")
            except Exception as e:
                logger.debug(f"[팝업워치독] WM_CLOSE 실패: {e}")

        with self._lock:
            self.dismiss_count += 1
            count = self.dismiss_count
            granted_now = False
            if count >= self.rest_threshold:
                if self._pending_rest == 0:
                    self._pending_rest = 1
                    self.rest_granted_total += 1
                    granted_now = True
            pending = self._pending_rest

        try:
            if _overlay is not None:
                if pending > 0:
                    _overlay.set_status(
                        f"⚠ OS 팝업 {count}회 — 휴식 {self.rest_seconds}초 예정"
                    )
                else:
                    _overlay.set_status(f"⚠ OS 팝업 {count}회 (자동 닫음)")
        except Exception:
            pass

    def _click_button(self, user32, parent_hwnd):
        try:
            child = user32.GetWindow(parent_hwnd, self._GW_CHILD)
            seen = 0
            while child and seen < 30:
                seen += 1
                t = self._get_window_text(user32, child)
                if t and any(b in t for b in self._BUTTON_TEXTS):
                    try:
                        user32.PostMessageW(child, self._BM_CLICK, 0, 0)
                        logger.info(f"[팝업워치독] 버튼 클릭: '{t}'")
                        return True
                    except Exception:
                        pass
                child = user32.GetWindow(child, self._GW_HWNDNEXT)
        except Exception:
            pass
        return False


# ═══════════════════════════════════════════════════════════
# 오버레이 헬퍼
# ═══════════════════════════════════════════════════════════
def overlay_status(text):
    try:
        if _overlay is not None:
            _overlay.set_status(text)
    except Exception:
        pass


def overlay_progress(current, total):
    try:
        if _overlay is not None:
            _overlay.set_progress(current, total)
    except Exception:
        pass


# ═══════════════════════════════════════════════════════════
# 절전/화면보호기 방지 (Windows 전용)
# ═══════════════════════════════════════════════════════════
_ES_CONTINUOUS        = 0x80000000
_ES_SYSTEM_REQUIRED   = 0x00000001
_ES_DISPLAY_REQUIRED  = 0x00000002


def prevent_sleep():
    if platform.system() != "Windows":
        logger.info("절전 방지 건너뜀 (Windows 아님)")
        return False
    try:
        result = ctypes.windll.kernel32.SetThreadExecutionState(
            _ES_CONTINUOUS | _ES_SYSTEM_REQUIRED | _ES_DISPLAY_REQUIRED
        )
        if result == 0:
            logger.warning("절전 방지 요청이 거부됨")
            return False
        logger.info("절전/화면 꺼짐/화면보호기 방지 활성화")
        return True
    except Exception as e:
        logger.warning(f"절전 방지 설정 실패 (무시하고 계속): {e}")
        return False


def allow_sleep():
    if platform.system() != "Windows":
        return
    try:
        ctypes.windll.kernel32.SetThreadExecutionState(_ES_CONTINUOUS)
        logger.info("절전/화면보호기 방지 해제")
    except Exception as e:
        logger.debug(f"절전 방지 해제 실패 (무시): {e}")


# ═══════════════════════════════════════════════════════════
# 주소 파서 — "[번지 앞] + 마지막 N동 N호" 질의어 생성
# ═══════════════════════════════════════════════════════════
def _before_jeon(addr, jeon_list, cutoff):
    """jeon_list 중 end 위치가 cutoff 이하인 가장 마지막 '번지' 앞 텍스트 반환.
    해당하는 것이 없으면 None.
    """
    result = None
    for m in jeon_list:
        if m.end() <= cutoff:
            result = addr[:m.start()].strip()
    return result


def parse_address_to_query(address):
    """주소 문자열을 메인 검색창에 입력할 질의어로 변환.

    우선순위:
      1. "번지" 있고 "N동 N호" 있으면  → "번지 앞" + "N동 N호"
      2. "번지" 있고 "N호" 만 있으면   → "번지 앞" + "N호"  (동 생략 케이스)
      3. "번지" 있고 동호수 없으면      → "번지 앞" 만
      4. "번지" 없으면                  → 전체 주소에서 "아파트" 또는 "마을"
                                          포함 토큰 제거 후 나머지

    예:
      "경기도 고양시 일산동구 마두동 732번지 백마한양아파트 310동 1005호"
        → "경기도 고양시 일산동구 마두동 732 310동 1005호"
      "경기도 고양시 일산동구 마두동 732번지 백마한양아파트 1005호"  (동 생략)
        → "경기도 고양시 일산동구 마두동 732 1005호"
      "경기도 고양시 일산동구 마두동 732번지"  (동호수 없음)
        → "경기도 고양시 일산동구 마두동 732"
      "경기도 고양시 일산동구 일산로 206 백마마을아파트 310동 1005호"  (번지 없음)
        → "경기도 고양시 일산동구 일산로 206 310동 1005호"

    Returns:
      (query: str | None, dong_ho_found: bool)
    """
    if not address:
        return None, False

    addr = re.sub(r"\s+", " ", str(address).strip())

    dong_ho_list = list(re.finditer(r"(\d+)\s*동\s*(\d+)\s*호", addr))
    ho_only_list = list(re.finditer(r"(\d+)\s*호", addr))
    jeon_list    = list(re.finditer(r"번지", addr))

    # ── 경우 4: 번지 없음 ────────────────────────────────────────
    if not jeon_list:
        tokens   = addr.split()
        filtered = [t for t in tokens if "아파트" not in t and "마을" not in t]
        query    = re.sub(r"\s+", " ", " ".join(filtered)).strip()
        return (query or None), bool(dong_ho_list)

    # ── 경우 1: N동 N호 있음 ─────────────────────────────────────
    if dong_ho_list:
        last_dh = dong_ho_list[-1]
        before  = _before_jeon(addr, jeon_list, cutoff=last_dh.start())
        if before is None:
            before = addr[:last_dh.start()].strip()
        dong, ho = last_dh.group(1), last_dh.group(2)
        query = re.sub(r"\s+", " ", f"{before} {dong}동 {ho}호").strip()
        return (query or None), True

    # ── 경우 2: N호만 있음 (동 생략) ─────────────────────────────
    if ho_only_list:
        last_ho = ho_only_list[-1]
        before  = _before_jeon(addr, jeon_list, cutoff=last_ho.start())
        if before is None:
            before = addr[:last_ho.start()].strip()
        ho    = last_ho.group(1)
        query = re.sub(r"\s+", " ", f"{before} {ho}호").strip()
        return (query or None), False

    # ── 경우 3: 동호수 모두 없음 ─────────────────────────────────
    last_jeon = jeon_list[-1]
    query     = re.sub(r"\s+", " ", addr[:last_jeon.start()].strip()).strip()
    return (query or None), False


# ═══════════════════════════════════════════════════════════
# 엑셀 입출력
# ═══════════════════════════════════════════════════════════
def _norm_header(value):
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value)).lower()


def _match_header(value, keyword_list):
    norm = _norm_header(value)
    if not norm:
        return False
    for kw in keyword_list:
        if _norm_header(kw) in norm:
            return True
    return False


def detect_columns(ws):
    """워크시트 1행 헤더를 스캔해 컬럼 인덱스 매핑을 반환.

    Returns:
      dict with keys: address, task_no, unique, kind, repr, note
                       각 값은 1-based 컬럼 인덱스 또는 None.
      또한 header_row 키는 헤더 행 번호 (기본 1).
    """
    header_row = 1
    max_col = ws.max_column or 1

    cols = {
        "address": None,
        "task_no": None,
        "unique":  None,
        "kind":    None,
        "repr":    None,
        "note":    None,
    }

    headers_found = {}
    for c in range(1, max_col + 1):
        v = ws.cell(row=header_row, column=c).value
        if v is None:
            continue
        headers_found[c] = v

    # 정확한 단어 우선 매칭. 더 구체적인 헤더(부동산고유번호, 부동산구분, 부동산표시)
    # 부터 매칭하여 일반 키워드("번호") 가 과도하게 잡히는 것을 방지.
    # 1) 부동산고유번호
    for c, v in headers_found.items():
        if _match_header(v, HEADER_UNIQUE_NO_KEYWORDS):
            cols["unique"] = c
            break
    # 2) 부동산구분
    for c, v in headers_found.items():
        if c == cols["unique"]:
            continue
        if _match_header(v, HEADER_KIND_KEYWORDS):
            cols["kind"] = c
            break
    # 3) 부동산표시
    for c, v in headers_found.items():
        if c in (cols["unique"], cols["kind"]):
            continue
        if _match_header(v, HEADER_REPR_KEYWORDS):
            cols["repr"] = c
            break
    # 4) 비고
    for c, v in headers_found.items():
        if c in (cols["unique"], cols["kind"], cols["repr"]):
            continue
        if _match_header(v, HEADER_NOTE_KEYWORDS):
            cols["note"] = c
            break
    # 5) 주소
    for c, v in headers_found.items():
        if _match_header(v, HEADER_ADDRESS_KEYWORDS):
            cols["address"] = c
            break
    # 6) 작업번호 — 출력 컬럼/주소 컬럼과 겹치지 않게.
    busy = {cols["unique"], cols["kind"], cols["repr"], cols["note"], cols["address"]}
    for c, v in headers_found.items():
        if c in busy:
            continue
        if _match_header(v, HEADER_TASK_NO_KEYWORDS):
            cols["task_no"] = c
            break

    cols["header_row"] = header_row
    return cols


def ensure_output_columns(ws, cols):
    """결과 헤더(부동산고유번호/부동산구분/부동산표시/비고) 가 없으면
    주소 컬럼 우측에 4개를 새로 추가하고 cols 매핑을 갱신한다.
    헤더는 1행에 작성.
    """
    if cols["address"] is None:
        return False  # 주소 컬럼이 없으면 처리 불가

    # 어느 하나라도 결과 컬럼이 없으면, 주소 컬럼 우측부터 차례로 보강.
    next_col = ws.max_column + 1
    # 우측 첫 빈 컬럼부터 추가하되, 주소 직후가 비어있으면 그쪽을 우선.
    candidate = cols["address"] + 1
    while candidate <= ws.max_column:
        if ws.cell(row=1, column=candidate).value is None:
            next_col = candidate
            break
        candidate += 1
    if candidate > ws.max_column:
        next_col = ws.max_column + 1

    for key, label in (
        ("unique", DEFAULT_HEADER_UNIQUE),
        ("kind",   DEFAULT_HEADER_KIND),
        ("repr",   DEFAULT_HEADER_REPR),
        ("note",   DEFAULT_HEADER_NOTE),
    ):
        if cols[key] is None:
            ws.cell(row=1, column=next_col, value=label)
            cols[key] = next_col
            next_col += 1

    return True


def load_address_rows(file_path):
    """엑셀에서 (행번호, 작업번호, 주소) 목록을 읽어 반환.

    주소 컬럼이 비어있는 행은 자동 스킵.
    Returns:
      (items, cols)
        items: list of dict {row, task_no, address}
        cols: 컬럼 매핑 dict (output 컬럼 보장됨)
    """
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    cols = detect_columns(ws)
    if cols["address"] is None:
        wb.close()
        raise ValueError(
            "엑셀 1행 헤더에서 '주소' 컬럼을 찾지 못했습니다. "
            "헤더에 '주소' 라는 단어가 포함된 컬럼을 추가해 주세요."
        )

    ok = ensure_output_columns(ws, cols)
    if not ok:
        wb.close()
        raise ValueError("결과 출력 컬럼 보강 실패")

    # 변경사항(헤더 보강) 저장.
    wb.save(file_path)

    items = []
    for row in range(cols["header_row"] + 1, ws.max_row + 1):
        addr_v = ws.cell(row=row, column=cols["address"]).value
        if addr_v is None or str(addr_v).strip() == "":
            continue
        task_v = None
        if cols.get("task_no"):
            task_v = ws.cell(row=row, column=cols["task_no"]).value
        items.append({
            "row": row,
            "task_no": ("" if task_v is None else str(task_v).strip()),
            "address": str(addr_v).strip(),
        })

    wb.close()
    return items, cols


def write_result_row(file_path, cols, row, unique=None, kind=None, repr_=None, note=None):
    """엑셀의 특정 행에 결과를 기록 (즉시 저장).

    같은 행 주소 셀의 글꼴·정렬 서식을 복사하여 '주변 서식에 맞추기' 적용.
    """
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    # 참조 서식: 같은 행 주소 셀
    ref_font = ref_align = None
    if cols.get("address"):
        try:
            ref = ws.cell(row=row, column=cols["address"])
            if ref.has_style:
                ref_font  = copy(ref.font)
                ref_align = copy(ref.alignment)
        except Exception:
            pass

    def _write(col, value):
        if col is None or value is None:
            return
        cell = ws.cell(row=row, column=col, value=value)
        if ref_font is not None:
            try:
                cell.font = copy(ref_font)
            except Exception:
                pass
        if ref_align is not None:
            try:
                cell.alignment = copy(ref_align)
            except Exception:
                pass

    _write(cols.get("unique"), unique)
    _write(cols.get("kind"),   kind)
    _write(cols.get("repr"),   repr_)
    _write(cols.get("note"),   note)

    # 잠금된 경우 대비 짧은 재시도
    for attempt in range(3):
        try:
            wb.save(file_path)
            break
        except PermissionError:
            logger.warning(
                f"엑셀 파일 저장 잠금 (시도 {attempt+1}/3) — 2초 후 재시도. "
                f"파일을 다른 프로그램에서 열고 있다면 닫아주세요: {file_path}"
            )
            time.sleep(2)
    else:
        logger.error(f"엑셀 파일 저장 실패 (잠김): {file_path}")
    wb.close()


# ═══════════════════════════════════════════════════════════
# 인터넷등기소 자동화 클래스
# ═══════════════════════════════════════════════════════════
class PropertyIdCollector:

    def __init__(self):
        self.driver = None

    # ── 유틸 ──────────────────────────────────────────────

    def _dismiss_any_alert(self):
        try:
            alert = self.driver.switch_to.alert
            alert.accept()
            time.sleep(0.5)
            return True
        except NoAlertPresentException:
            return False
        except Exception:
            return False

    def _dismiss_modal_popup(self):
        try:
            removed = self.driver.execute_script("""
                var modal = document.getElementById('_modal');
                if (modal && modal.style.display === 'block') {
                    modal.style.display = 'none';
                    return true;
                }
                var modals = document.querySelectorAll('div.w2modal_popup[style*="display: block"]');
                var found = false;
                for (var i = 0; i < modals.length; i++) {
                    modals[i].style.display = 'none';
                    found = true;
                }
                return found;
            """)
            if removed:
                logger.info("잔존 모달 팝업 강제 제거")
                time.sleep(0.3)
            return removed
        except Exception:
            return False

    def _kill_processbar_overlay(self):
        try:
            killed = self.driver.execute_script("""
                var bars = document.querySelectorAll(
                    'div[id*="processbar"], div.w2modal'
                );
                var cnt = 0;
                for (var i = 0; i < bars.length; i++) {
                    var s = window.getComputedStyle(bars[i]);
                    if (s.display !== 'none' && s.visibility !== 'hidden') {
                        bars[i].style.display = 'none';
                        cnt++;
                    }
                }
                return cnt;
            """)
            if killed:
                logger.warning(f"processbar 오버레이 {killed}개 강제 제거")
                time.sleep(0.3)
            return killed or 0
        except Exception as e:
            logger.debug(f"processbar 제거 중 오류 (무시): {e}")
            return 0

    def _wait_until_processbar_gone(self, max_wait=20.0, poll=0.2):
        start = time.time()
        deadline = start + max_wait
        js = """
            var bars = document.querySelectorAll('div[id*="processbar"]');
            for (var i = 0; i < bars.length; i++) {
                var s = window.getComputedStyle(bars[i]);
                if (s.display !== 'none' && s.visibility !== 'hidden') {
                    return true;
                }
            }
            return false;
        """
        while time.time() < deadline:
            try:
                visible = self.driver.execute_script(js)
            except Exception:
                visible = False
            if not visible:
                waited = time.time() - start
                if waited > 0.5:
                    logger.info(f"[처리중 바] 사라짐 (대기 {waited:.1f}초)")
                return True
            time.sleep(poll)
        logger.warning(f"[처리중 바] {max_wait:.0f}초 후에도 사라지지 않음")
        return False

    def _click_with_processbar_retry(self, element, label="click"):
        try:
            element.click()
            return True
        except ElementClickInterceptedException:
            logger.warning(f"[{label}] click intercepted → processbar 제거 후 재시도")
            killed = self._kill_processbar_overlay()
            if killed > 0:
                try:
                    element.click()
                    return True
                except Exception:
                    pass
            try:
                self.driver.execute_script("arguments[0].click();", element)
                logger.info(f"[{label}] JS click 폴백 성공")
                return True
            except Exception as e2:
                logger.error(f"[{label}] 모든 클릭 경로 실패: {e2}")
                return False
        except Exception as e:
            logger.debug(f"[{label}] click 예외(intercept 아님): {e}")
            raise

    def _handle_alert(self, accept=True, timeout=1):
        try:
            alert = WebDriverWait(self.driver, timeout).until(EC.alert_is_present())
            text = alert.text
            logger.info(f"알림: {text[:80]}")
            alert.accept() if accept else alert.dismiss()
            time.sleep(0.5)
            return text
        except TimeoutException:
            return None

    def _get_body_text(self):
        try:
            return self.driver.execute_script(
                "return (document.body && (document.body.innerText || document.body.textContent)) || '';"
            ) or ""
        except Exception:
            try:
                return self.driver.find_element(By.TAG_NAME, "body").text or ""
            except Exception:
                return ""

    def _safe_click(self, elem):
        try:
            self.driver.execute_script(
                "arguments[0].scrollIntoView({block:'center', inline:'center'});", elem)
        except Exception:
            pass
        for action in (
            lambda: elem.click(),
            lambda: ActionChains(self.driver).move_to_element(elem).pause(0.2).click(elem).perform(),
            lambda: self.driver.execute_script("arguments[0].click();", elem),
        ):
            try:
                action()
                return True
            except (ElementClickInterceptedException,
                    ElementNotInteractableException,
                    StaleElementReferenceException):
                continue
            except Exception:
                continue
        return False

    def _click_any_visible(self, selectors, timeout=3, label=None):
        for by, sel in selectors:
            try:
                elems = WebDriverWait(self.driver, timeout).until(
                    EC.presence_of_all_elements_located((by, sel))
                )
            except TimeoutException:
                continue
            for elem in elems:
                try:
                    if not elem.is_displayed():
                        continue
                    if self._safe_click(elem):
                        if label:
                            logger.info(f"{label} 클릭")
                        time.sleep(0.5)
                        return True
                except StaleElementReferenceException:
                    continue
                except Exception:
                    continue
        return False

    # ── 세션 만료/페이지 진단 ──────────────────────────────

    def _is_page_blank(self):
        try:
            text_len = self.driver.execute_script(
                "return (document.body && document.body.innerText || '').trim().length;")
            return text_len < 30
        except Exception:
            return True

    def _wait_for_page_content(self, max_retries=10, blank_threshold_sec=10.0):
        """빈 화면이 blank_threshold_sec 이상 지속될 때만 새로고침 (최대 max_retries 회).

        document.readyState 가 complete 가 되기를 먼저 기다리고,
        그래도 본문이 비어있는 상태가 blank_threshold_sec 초 동안 계속되면 새로고침.
        """
        for attempt in range(max_retries):
            try:
                WebDriverWait(self.driver, 10).until(
                    lambda d: d.execute_script("return document.readyState") == "complete"
                )
            except Exception:
                pass
            self._dismiss_any_alert()
            if not self._is_page_blank():
                return True
            # 빈 화면이지만 곧 채워질 수 있으므로 threshold 동안 폴링.
            deadline = time.time() + blank_threshold_sec
            filled = False
            while time.time() < deadline:
                time.sleep(0.5)
                self._dismiss_any_alert()
                if not self._is_page_blank():
                    filled = True
                    break
            if filled:
                return True
            logger.warning(
                f"빈 화면 {blank_threshold_sec:.0f}초 지속 → 새로고침 "
                f"(시도 {attempt+1}/{max_retries})"
            )
            try:
                self.driver.refresh()
            except Exception:
                self.driver.get(IROS_INDEX_URL)
        return False

    def _is_security_install_page(self):
        try:
            title = self.driver.title or ""
            if "TouchEn" in title or "보안" in title or "제품 설치" in title:
                return True
            url = self.driver.current_url or ""
            if "touchen" in url.lower() or "nxkey" in url.lower():
                return True
            body = self._get_body_text()[:500]
            if "보안 프로그램 설치" in body or "TouchEn" in body:
                return True
        except Exception:
            pass
        return False

    def _recover_from_security_page(self):
        logger.warning("보안 프로그램 설치 페이지 감지 → 인터넷등기소로 복구")
        try:
            self.driver.get(IROS_INDEX_URL)
            self._wait_for_page_content()
            if self._is_security_install_page():
                time.sleep(2)
                self.driver.get(IROS_INDEX_URL)
                self._wait_for_page_content()
            return not self._is_security_install_page()
        except Exception as e:
            logger.error(f"보안 페이지 복구 실패: {e}")
            return False

    def _is_session_expired(self):
        try:
            alert = self.driver.switch_to.alert
            alert_text = alert.text
            if any(kw in alert_text for kw in
                   ["세션", "만료", "시간", "로그인",
                    "연결이 끊어", "다시 로그인", "서버와 연결"]):
                logger.warning(f"세션 만료 alert: {alert_text[:80]}")
                alert.accept()
                time.sleep(2)
                return True
            alert.accept()
            time.sleep(1)
        except NoAlertPresentException:
            pass

        try:
            session_keywords = ["세션", "만료", "자동 로그아웃", "자동로그아웃",
                                "연결이 끊어", "다시 로그인", "서버와 연결"]
            for kw in session_keywords:
                elems = self.driver.find_elements(
                    By.XPATH, f"//*[contains(text(),'{kw}')]")
                for elem in elems:
                    try:
                        if not elem.is_displayed():
                            continue
                        logger.warning(f"세션 만료 팝업 감지: {elem.text[:60]}")
                        popup_root = elem
                        for _ in range(6):
                            try:
                                popup_root = popup_root.find_element(By.XPATH, "..")
                                btns = popup_root.find_elements(
                                    By.XPATH,
                                    ".//input[@type='button'] | .//button | .//a[contains(@class,'btn')]")
                                for btn in btns:
                                    if btn.is_displayed():
                                        try:
                                            btn.click()
                                        except Exception:
                                            self.driver.execute_script("arguments[0].click();", btn)
                                        time.sleep(2)
                                        logger.info("세션 만료 팝업 버튼 클릭")
                                        return True
                            except Exception:
                                break
                    except Exception:
                        continue
        except Exception:
            pass

        # 헤더 로그아웃 그룹 visibility 체크
        try:
            grp_login = self.driver.find_element(
                By.ID, "mf_wfm_potal_main_wf_header_grp_login")
            grp_logout = self.driver.find_element(
                By.ID, "mf_wfm_potal_main_wf_header_grp_logout")
            login_style = grp_login.get_attribute("style") or ""
            logout_style = grp_logout.get_attribute("style") or ""
            if ("display: none" not in login_style and "visibility: hidden" not in login_style and
                    ("display: none" in logout_style or "visibility: hidden" in logout_style)):
                return True
        except Exception:
            pass

        try:
            if "로그인" in (self.driver.title or ""):
                return True
        except Exception:
            pass

        return False

    def ensure_logged_in(self):
        if self._is_security_install_page():
            self._recover_from_security_page()
        if self._is_session_expired():
            logger.warning("세션 만료 → 재로그인")
            self._dismiss_any_alert()
            time.sleep(2)
            self._wait_for_page_content()
            self._dismiss_modal_popup()
            self.login()
        return True

    # ── ChromeDriver / 프로필 ─────────────────────────────

    def _resolve_chromedriver_path(self):
        try:
            try:
                mgr = ChromeDriverManager(ssl_verify=False)
            except TypeError:
                mgr = ChromeDriverManager()
            path = mgr.install()
            if path and os.path.isfile(path):
                logger.info(f"ChromeDriver: webdriver_manager 로 획득 ({path})")
                return path
        except Exception as e:
            logger.warning(
                f"webdriver_manager 획득 실패 ({type(e).__name__}): {str(e)[:120]}"
            )

        try:
            cache_root = Path.home() / ".wdm" / "drivers" / "chromedriver"
            if cache_root.is_dir():
                candidates = list(cache_root.rglob("chromedriver.exe")) \
                             + list(cache_root.rglob("chromedriver"))
                candidates = [c for c in candidates if c.is_file()]
                candidates.sort(key=lambda p: p.stat().st_mtime, reverse=True)
                if candidates:
                    path = str(candidates[0])
                    logger.info(f"ChromeDriver: 로컬 캐시 ({path})")
                    return path
        except Exception as e:
            logger.warning(f"로컬 캐시 탐색 실패: {e}")

        logger.warning("ChromeDriver: PATH 에서 탐색 시도")
        return None

    def _prepare_selenium_profile(self):
        """기존 Chrome 프로필에서 보안 모듈 관련 데이터만 복사한 경량 프로필."""
        chrome_src = Path.home() / "AppData" / "Local" / "Google" / "Chrome" / "User Data"
        if not chrome_src.is_dir():
            logger.info("Chrome 프로필 경로 없음 → 기본 프로필 사용")
            return None

        selenium_profile = Path.home() / ".selenium_chrome_profile"
        default_src = chrome_src / "Default"
        default_dst = selenium_profile / "Default"

        try:
            default_dst.mkdir(parents=True, exist_ok=True)

            copy_targets = [
                "Local Storage",
                "IndexedDB",
                "Extension State",
                "Extensions",
                "Preferences",
                "Secure Preferences",
            ]
            for target in copy_targets:
                src = default_src / target
                dst = default_dst / target
                if not src.exists():
                    continue
                try:
                    if src.is_dir():
                        if dst.exists():
                            shutil.rmtree(str(dst), ignore_errors=True)
                        shutil.copytree(str(src), str(dst), dirs_exist_ok=True)
                    else:
                        shutil.copy2(str(src), str(dst))
                except Exception:
                    pass

            for fname in ["Local State"]:
                src_f = chrome_src / fname
                if src_f.is_file():
                    try:
                        shutil.copy2(str(src_f), str(selenium_profile / fname))
                    except Exception:
                        pass

            logger.info(f"경량 Chrome 프로필 생성 완료: {selenium_profile}")
            return selenium_profile
        except Exception as e:
            logger.warning(f"경량 프로필 생성 실패: {e}")
            return None

    def start(self):
        logger.info("브라우저 시작 중...")
        opts = Options()
        opts.add_argument("--no-sandbox")
        opts.add_argument("--disable-dev-shm-usage")
        opts.add_argument("--disable-gpu")
        opts.add_argument("--start-maximized")
        opts.add_experimental_option("excludeSwitches", ["enable-automation"])
        opts.add_experimental_option("useAutomationExtension", False)

        profile_dir = self._prepare_selenium_profile()
        if profile_dir:
            opts.add_argument(f"--user-data-dir={profile_dir}")
            opts.add_argument("--profile-directory=Default")

        driver_path = self._resolve_chromedriver_path()
        service = Service(driver_path) if driver_path else Service()
        self.driver = webdriver.Chrome(service=service, options=opts)
        self.driver.set_page_load_timeout(30)
        self.driver.get(IROS_INDEX_URL)
        self._wait_for_page_content()
        logger.info("인터넷등기소 접속 완료")

    # ── 로그인 ────────────────────────────────────────────

    def _login_form_visible(self):
        try:
            el = self.driver.find_element(
                By.ID, "mf_wfm_potal_main_wfm_content_sbx_user_id_g___input")
            return self.driver.execute_script(
                "return arguments[0].offsetParent !== null;", el)
        except Exception:
            return False

    def _is_logged_in(self):
        """현재 로그인 상태 확인. 헤더의 로그아웃 그룹 visibility 기준."""
        try:
            grp_logout = self.driver.find_element(
                By.ID, "mf_wfm_potal_main_wf_header_grp_logout")
            style = grp_logout.get_attribute("style") or ""
            if "display: none" not in style and "visibility: hidden" not in style:
                return True
        except Exception:
            pass
        try:
            txt = self.driver.execute_script(
                "var e=document.getElementById('mf_wfm_potal_main_wf_header_grp_logout');"
                "if(!e) return '';"
                "return (e.innerText||e.textContent||'');") or ""
            if "로그아웃" in txt:
                return True
        except Exception:
            pass
        return False

    def _try_open_login_from_current_page(self, label_prefix=""):
        self._dismiss_any_alert()

        top_login_selectors = [
            (By.ID, "mf_wfm_potal_main_wf_header_btn_login"),
            (By.XPATH, "//*[@id='mf_wfm_potal_main_wf_header_grp_login']//*[self::a or self::button][contains(normalize-space(.), '로그인')]"),
            (By.XPATH, "//*[contains(@id,'grp_login')]//*[self::a or self::button][contains(normalize-space(.), '로그인')]"),
            (By.XPATH, "//a[contains(normalize-space(.), '로그인')][ancestor::*[contains(@id,'header')]]"),
        ]
        dropdown_login_selectors = [
            (By.ID, "mf_wfm_potal_main_wf_header_wq_uuid_2529"),
            (By.XPATH, "//*[@id='mf_wfm_potal_main_wf_header_grp_login']//div[contains(@class,'box-link')]//*[self::a or self::button][contains(normalize-space(.), '로그인')]"),
            (By.XPATH, "//*[contains(@class,'box-link')]//*[self::a or self::button][contains(normalize-space(.), '로그인')]"),
        ]

        if self._click_any_visible(top_login_selectors, timeout=2, label=f"{label_prefix}상단 로그인"):
            try:
                WebDriverWait(self.driver, 2).until(
                    lambda d: d.execute_script(
                        "var box = document.querySelector('#mf_wfm_potal_main_wf_header_grp_login .box-link');"
                        "return !!(box && getComputedStyle(box).display !== 'none' && getComputedStyle(box).visibility !== 'hidden');"
                    )
                )
            except TimeoutException:
                pass

        try:
            self.driver.execute_script("""
                var grp = document.getElementById('mf_wfm_potal_main_wf_header_grp_login');
                if (grp) { grp.classList.add('on'); }
                var box = document.querySelector('#mf_wfm_potal_main_wf_header_grp_login .box-link');
                if (box) {
                    box.style.display = 'block';
                    box.style.visibility = 'visible';
                    box.style.opacity = '1';
                    box.style.pointerEvents = 'auto';
                }
            """)
        except Exception:
            pass

        if self._click_any_visible(dropdown_login_selectors, timeout=2, label=f"{label_prefix}드롭다운 로그인"):
            try:
                WebDriverWait(self.driver, 5).until(lambda d: self._login_form_visible())
                return True
            except TimeoutException:
                pass

        return self._login_form_visible()

    def _open_login_form(self):
        if self._login_form_visible():
            return True

        self._dismiss_any_alert()

        try:
            if self._try_open_login_from_current_page():
                return True
        except Exception as e:
            logger.warning(f"현재 화면에서 로그인 진입 실패: {e}")

        try:
            current_url = self.driver.current_url or ""
        except Exception:
            current_url = ""

        if IROS_BASE_URL not in current_url:
            try:
                self.driver.get(IROS_INDEX_URL)
                WebDriverWait(self.driver, 8).until(
                    lambda d: d.execute_script("return document.readyState") in ("interactive", "complete")
                )
                self._dismiss_any_alert()
                if self._try_open_login_from_current_page(label_prefix="메인 재진입 후 "):
                    return True
            except Exception as e:
                logger.warning(f"메인 재진입 후 로그인 실패: {e}")

        try:
            self.driver.execute_script("""
                var grp = document.getElementById('mf_wfm_potal_main_wf_header_grp_login');
                if (grp) { grp.classList.add('on'); }
                var box = document.querySelector('#mf_wfm_potal_main_wf_header_grp_login .box-link');
                if (box) {
                    box.style.display = 'block';
                    box.style.visibility = 'visible';
                    box.style.opacity = '1';
                    box.style.pointerEvents = 'auto';
                }
                var explicit = document.getElementById('mf_wfm_potal_main_wf_header_wq_uuid_2529');
                if (explicit) {
                    explicit.dispatchEvent(new MouseEvent('click', {bubbles:true, cancelable:true, view:window}));
                    explicit.click();
                    return true;
                }
                var links = document.querySelectorAll('#mf_wfm_potal_main_wf_header_grp_login a, #mf_wfm_potal_main_wf_header_grp_login button');
                for (var i = 0; i < links.length; i++) {
                    var txt = (links[i].innerText || links[i].textContent || '').trim();
                    if (txt === '로그인' || txt.indexOf('로그인') >= 0) {
                        links[i].dispatchEvent(new MouseEvent('click', {bubbles:true, cancelable:true, view:window}));
                        links[i].click();
                        return true;
                    }
                }
                return false;
            """)
            WebDriverWait(self.driver, 5).until(lambda d: self._login_form_visible())
            logger.info("JS fallback 으로 로그인 폼 로드 완료")
            return True
        except Exception as e:
            logger.warning(f"JS fallback 로그인 진입 실패: {e}")

        return self._login_form_visible()

    def login(self):
        self._dismiss_any_alert()

        if self._is_logged_in():
            logger.info("이미 로그인 상태")
            return True

        logger.info("로그인 시도...")
        if not self._open_login_form():
            logger.error("로그인 페이지 진입 실패")
            return False

        self._dismiss_modal_popup()
        self._dismiss_any_alert()

        for login_attempt in range(2):
            try:
                wait = WebDriverWait(self.driver, ELEMENT_WAIT_TIMEOUT)

                id_field = wait.until(EC.element_to_be_clickable(
                    (By.ID, "mf_wfm_potal_main_wfm_content_sbx_user_id_g___input")))
                pw_field = wait.until(EC.element_to_be_clickable(
                    (By.ID, "mf_wfm_potal_main_wfm_content_sct_mbr_pw_g")))

                id_field.click()
                time.sleep(0.2)
                id_field.clear()
                id_field.send_keys(LOGIN_ID)
                time.sleep(0.2)

                actual_id = id_field.get_attribute("value") or ""
                if actual_id != LOGIN_ID:
                    logger.warning(f"ID 입력값 불일치: '{actual_id}' → JS 재입력")
                    self.driver.execute_script(
                        "var e=arguments[0]; e.value=''; e.value=arguments[1];"
                        "e.dispatchEvent(new Event('input',{bubbles:true}));"
                        "e.dispatchEvent(new Event('change',{bubbles:true}));",
                        id_field, LOGIN_ID)
                    time.sleep(0.2)

                pw_field.click()
                time.sleep(0.2)
                pw_field.clear()
                pw_field.send_keys(LOGIN_PW)
                time.sleep(0.2)

                actual_pw = pw_field.get_attribute("value") or ""
                if not actual_pw:
                    logger.warning("PW send_keys 실패 → JS 직접 입력")
                    self.driver.execute_script(
                        "var e=arguments[0]; e.focus(); e.value=''; e.value=arguments[1];"
                        "e.dispatchEvent(new Event('input',{bubbles:true}));"
                        "e.dispatchEvent(new Event('change',{bubbles:true}));",
                        pw_field, LOGIN_PW)
                    time.sleep(0.3)
                    actual_pw = pw_field.get_attribute("value") or ""

                if not actual_pw:
                    logger.error(f"PW 입력 실패 (시도 {login_attempt+1}/2)")
                    if login_attempt == 0:
                        self._dismiss_modal_popup()
                        self._dismiss_any_alert()
                        time.sleep(1)
                        continue
                    return False

                logger.info("ID/PW 입력 완료")

                login_btn = wait.until(EC.element_to_be_clickable(
                    (By.ID, "mf_wfm_potal_main_wfm_content_btn_login")))
                self._safe_click(login_btn)
                time.sleep(3)
                self._dismiss_any_alert()

                for _ in range(8):
                    if self._is_logged_in():
                        logger.info("로그인 완료")
                        return True
                    try:
                        current_title = self.driver.title or ""
                        current_url = self.driver.current_url or ""
                        if (not self._login_form_visible() and "로그인" not in current_title
                                and "POPM10_P00_01" not in current_url):
                            logger.info("로그인 완료(폼 종료 기준)")
                            return True
                    except Exception:
                        pass
                    time.sleep(1)

                alert_text = self._handle_alert(accept=True)
                if alert_text and ("패스워드" in alert_text or "비밀번호" in alert_text):
                    logger.warning(f"비밀번호 미입력 알림: {alert_text[:60]} → 재시도")
                    if login_attempt == 0:
                        self._dismiss_modal_popup()
                        time.sleep(1)
                        continue
                    return False

                logger.error("로그인 후에도 로그인 화면에서 벗어나지 못함")
                if login_attempt == 0:
                    self._dismiss_modal_popup()
                    self._dismiss_any_alert()
                    time.sleep(1)
                    continue
                return False

            except Exception as e:
                logger.error(f"로그인 실패: {e}")
                if login_attempt == 0:
                    self._dismiss_modal_popup()
                    self._dismiss_any_alert()
                    time.sleep(1)
                    continue
                return False

        return False

    # ── 메인 페이지 이동 ──────────────────────────────────

    def _is_truly_on_main(self):
        """검색창이 보이면서 동시에 결과 페이지가 아닌 경우에만 True.

        검색결과 페이지에도 상단에 검색창이 있어 단순 검색창 존재만으로는
        메인을 판정할 수 없다. 결과 페이지 고유 신호가 본문에 있으면 False.
        """
        if self._wait_main_search_box(timeout=0.5) is None:
            return False
        try:
            body = self._get_body_text()
            for kw in self._NO_RESULT_KEYWORDS:
                if kw in body:
                    return False
            if re.search(r"전체\s*\d+\s*건", body):
                return False
            if "부동산 등기사항증명서 열람·발급 신청" in body:
                return False
        except Exception:
            pass
        return True

    def go_to_main(self):
        """필요할 때만 메인 페이지로 이동.

        이미 메인이고 검색창이 보이면 아무 동작도 하지 않고 즉시 리턴.
        보안 페이지/세션 만료 등 비정상 상태일 때만 이동을 수행한다.
        """
        self._dismiss_any_alert()

        # 보안 설치 페이지 같은 비정상 상태부터 복구
        if self._is_security_install_page():
            self._recover_from_security_page()

        # 이미 진짜 메인이라면 그대로 진행
        if self._is_truly_on_main():
            logger.debug("메인 페이지 이미 표시 중 — 이동 생략")
            return True

        # 진짜로 이동이 필요한 경우에만 로고 클릭 → URL get
        try:
            logo = self.driver.find_element(By.ID, "mf_wfm_potal_main_wf_header_btn_home")
            self.driver.execute_script("arguments[0].click();", logo)
        except Exception:
            try:
                self.driver.get(IROS_INDEX_URL)
            except Exception:
                pass
        self._wait_for_page_content()
        self._wait_main_search_box(timeout=8)
        # 메인 도착 검증 — 실패 시 URL 강제 이동으로 한 번 더 보강
        if not self._is_truly_on_main():
            try:
                self.driver.get(IROS_INDEX_URL)
                self._wait_for_page_content()
                self._wait_main_search_box(timeout=8)
            except Exception:
                pass
        logger.info("메인 페이지 이동")
        return True
        return True

    # ── 메인 검색창 자동화 ────────────────────────────────

    # 메인 검색창 input 을 찾는 JS. 여러 단서로 견고하게 탐색.
    _FIND_MAIN_SEARCH_INPUT_JS = r"""
        // 1) placeholder 매칭 (가장 강력한 단서)
        var phHints = [
            '부동산 등기사항증명서',
            '주소를 입력',
            '주소를 입력하세요'
        ];
        var allInputs = document.querySelectorAll('input');
        for (var i = 0; i < allInputs.length; i++) {
            var el = allInputs[i];
            var s = window.getComputedStyle(el);
            if (s.display === 'none' || s.visibility === 'hidden') continue;
            if (el.offsetParent === null) continue;
            var ph = el.getAttribute('placeholder') || '';
            for (var k = 0; k < phHints.length; k++) {
                if (ph.indexOf(phHints[k]) !== -1) {
                    return el;
                }
            }
        }
        // 2) 메인 검색 영역 컨테이너 패턴
        var ctnHints = [
            'real_search', 'main_search', 'sbx_real', 'sbx_main', 'sbx_search'
        ];
        for (var i = 0; i < allInputs.length; i++) {
            var el = allInputs[i];
            var s = window.getComputedStyle(el);
            if (s.display === 'none' || s.visibility === 'hidden') continue;
            if (el.offsetParent === null) continue;
            var idn = (el.id || '') + ' ' + (el.name || '');
            idn = idn.toLowerCase();
            for (var k = 0; k < ctnHints.length; k++) {
                if (idn.indexOf(ctnHints[k]) !== -1) {
                    return el;
                }
            }
        }
        return null;
    """

    def _wait_main_search_box(self, timeout=10):
        deadline = time.time() + timeout
        while time.time() < deadline:
            try:
                el = self.driver.execute_script(
                    "return (function(){" + self._FIND_MAIN_SEARCH_INPUT_JS + "})();"
                )
                if el is not None:
                    return el
            except Exception:
                pass
            time.sleep(0.3)
        return None

    def _ensure_realty_tab_selected(self):
        """메인 검색창 좌측의 '부동산' 탭이 활성화되어 있는지 확인하고
        '법인' 탭으로 전환되어 있으면 '부동산' 으로 되돌린다.
        """
        try:
            self.driver.execute_script(r"""
                // 검색 입력창 근처에 '부동산' / '법인' 토글 버튼을 찾는다.
                var inp = (function(){__FIND__})();
                if (!inp) return false;
                var box = inp;
                for (var i = 0; i < 6 && box; i++) {
                    box = box.parentElement;
                    if (!box) break;
                    // 같은 그룹 안에서 '부동산' / '법인' 버튼 찾기
                    var btns = box.querySelectorAll('button, a, span, div');
                    var realty = null, corp = null;
                    for (var j = 0; j < btns.length; j++) {
                        var t = (btns[j].textContent || '').trim();
                        if (t === '부동산') realty = btns[j];
                        if (t === '법인')   corp   = btns[j];
                    }
                    if (realty && corp) {
                        // 부동산이 비활성처럼 보이면 클릭
                        var cls = (realty.className || '') + '';
                        if (cls.indexOf('active') === -1 && cls.indexOf('on') === -1) {
                            try { realty.click(); } catch(e) {}
                        }
                        return true;
                    }
                }
                return false;
            """.replace("__FIND__", self._FIND_MAIN_SEARCH_INPUT_JS))
        except Exception as e:
            logger.debug(f"부동산 탭 보장 중 예외 (무시): {e}")

    def search_address_on_main(self, query):
        """메인화면 부동산 검색창에 query 를 입력하고 검색 실행.

        Returns:
          bool — 검색 액션이 정상 트리거되었는지.
        """
        # 1) 메인 검색창 확보 — 메인 페이지가 아니면 메인으로 이동.
        inp_el = self._wait_main_search_box(timeout=5)
        if inp_el is None:
            self.go_to_main()
            inp_el = self._wait_main_search_box(timeout=10)
            if inp_el is None:
                logger.error("메인 검색창을 찾지 못함")
                return False

        # 2) 부동산 탭 활성화
        self._ensure_realty_tab_selected()

        # 3) WebElement 로 다시 잡기 (execute_script 가 반환한 element 는 신뢰)
        try:
            inp_id = self.driver.execute_script("return arguments[0].id || '';", inp_el)
        except Exception:
            inp_id = ""

        # 4) 클릭 → 비우기 → 입력
        try:
            try:
                self.driver.execute_script(
                    "arguments[0].scrollIntoView({block:'center', inline:'center'});", inp_el)
            except Exception:
                pass

            self._click_with_processbar_retry(inp_el, label="메인 검색창")
            time.sleep(0.2)

            # 값 비우기 (clear + JS 둘 다)
            try:
                inp_el.clear()
            except Exception:
                pass
            self.driver.execute_script(
                "arguments[0].value='';"
                "arguments[0].dispatchEvent(new Event('input',{bubbles:true}));"
                "arguments[0].dispatchEvent(new Event('change',{bubbles:true}));",
                inp_el)
            time.sleep(0.1)

            # 값 입력
            try:
                inp_el.send_keys(query)
            except Exception:
                self.driver.execute_script(
                    "arguments[0].value=arguments[1];"
                    "arguments[0].dispatchEvent(new Event('input',{bubbles:true}));"
                    "arguments[0].dispatchEvent(new Event('change',{bubbles:true}));",
                    inp_el, query)

            # 입력 결과 검증
            actual = ""
            try:
                actual = inp_el.get_attribute("value") or ""
            except Exception:
                pass
            if actual.strip() != query.strip():
                logger.warning(
                    f"검색창 값 불일치 (입력: '{actual}' / 기대: '{query}') → JS 재입력"
                )
                self.driver.execute_script(
                    "arguments[0].value=arguments[1];"
                    "arguments[0].dispatchEvent(new Event('input',{bubbles:true}));"
                    "arguments[0].dispatchEvent(new Event('change',{bubbles:true}));",
                    inp_el, query)
                time.sleep(0.2)

            logger.info(f"메인 검색창 입력: '{query}'")
        except Exception as e:
            logger.error(f"검색창 입력 실패: {e}")
            return False

        # 5) 검색 트리거: 돋보기 클릭 → Enter 키 → form.submit() 순으로 시도하되
        #    각 시도 후 URL/페이지 변경을 검증해 실효성 있는 트리거만 채택.
        triggers = [
            ("돋보기 아이콘", lambda: self._click_search_icon_near(inp_el)),
            ("Enter 키",     lambda: self._send_enter_via_dispatch(inp_el)),
            ("form.submit",  lambda: self._submit_search_form(inp_el)),
        ]

        for label, action in triggers:
            before = self._capture_trigger_state()
            try:
                called = action()
            except Exception as e:
                logger.debug(f"트리거 호출 예외({label}): {e}")
                continue
            if not called:
                logger.debug(f"트리거 미실행({label})")
                continue

            # alert (검색어 형식 오류 등) 우선 흡수 → 발생 시 즉시 실패 처리
            alert_text = self._handle_alert(accept=True, timeout=0.5)
            if alert_text:
                logger.warning(f"트리거({label}) 직후 alert: {alert_text[:80]}")
                return False

            ok, reason = self._verify_search_triggered(before, timeout=4.0)
            if ok:
                logger.info(f"검색 트리거 성공: {label} ({reason})")
                return True
            logger.warning(f"트리거({label}) 효과 없음 — 다음 방법 시도")

        logger.error("모든 검색 트리거 방법이 효과 없음")
        return False

    # ── 검색 트리거 보조 ────────────────────────────────────

    def _capture_trigger_state(self):
        """검색 트리거 전 스냅샷 — URL, 본문 길이, processbar 표시 여부."""
        try:
            url = self.driver.current_url or ""
        except Exception:
            url = ""
        try:
            body_len = self.driver.execute_script(
                "return (document.body && document.body.innerText || '').length;"
            ) or 0
        except Exception:
            body_len = 0
        try:
            has_pb = bool(self.driver.execute_script(r"""
                var bars = document.querySelectorAll('div[id*="processbar"]');
                for (var i = 0; i < bars.length; i++) {
                    var s = window.getComputedStyle(bars[i]);
                    if (s.display !== 'none' && s.visibility !== 'hidden') return true;
                }
                return false;
            """))
        except Exception:
            has_pb = False
        return {"url": url, "body_len": int(body_len), "has_pb": has_pb}

    # 결과 페이지 도달을 가리키는 키워드 (메인에는 존재 가능성 낮음).
    _RESULT_SIGNAL_KEYWORDS = (
        "입력하신 검색조건에 대한 결과가 없습니다",
        "부동산 등기사항증명서 열람·발급 신청",
        "부동산 소재지번 검색 결과",
    )

    def _verify_search_triggered(self, before, timeout=4.0, poll=0.2):
        """트리거가 실제 효과를 발휘했는지 확인.

        다음 중 하나라도 만족하면 성공:
          - URL 변경
          - processbar 신규 출현
          - 결과 페이지 고유 키워드 출현
          - 본문 텍스트 길이가 의미 있게 변화 (>= 200자 증감) + '전체 N건' 매치
        """
        deadline = time.time() + timeout
        while time.time() < deadline:
            try:
                url = self.driver.current_url or ""
            except Exception:
                url = before["url"]
            if url != before["url"]:
                return True, "URL 변경"

            try:
                has_pb = bool(self.driver.execute_script(r"""
                    var bars = document.querySelectorAll('div[id*="processbar"]');
                    for (var i = 0; i < bars.length; i++) {
                        var s = window.getComputedStyle(bars[i]);
                        if (s.display !== 'none' && s.visibility !== 'hidden') return true;
                    }
                    return false;
                """))
            except Exception:
                has_pb = before["has_pb"]
            if has_pb and not before["has_pb"]:
                return True, "processbar 출현"

            try:
                body = self._get_body_text()
            except Exception:
                body = ""
            for kw in self._RESULT_SIGNAL_KEYWORDS:
                if kw in body:
                    return True, f"결과 키워드('{kw[:10]}…') 출현"
            if re.search(r"전체\s*\d+\s*건", body):
                return True, "'전체 N건' 표시 출현"

            time.sleep(poll)
        return False, "변화 감지 안 됨"

    def _send_enter_via_dispatch(self, inp_el):
        """입력창에 포커스를 주고 Enter 키 이벤트를 dispatch.

        send_keys 가 WebSquare 의 keypress 핸들러를 안 자극하는 케이스가 있어
        keydown/keypress/keyup 세 이벤트를 모두 발화하는 폴백을 함께 시도.
        """
        try:
            inp_el.send_keys(Keys.RETURN)
        except Exception as e:
            logger.debug(f"send_keys(RETURN) 실패: {e}")

        try:
            self.driver.execute_script(r"""
                var el = arguments[0];
                if (!el) return false;
                try { el.focus(); } catch(e) {}
                var opts = {bubbles:true, cancelable:true, key:'Enter',
                            code:'Enter', keyCode:13, which:13, charCode:13};
                ['keydown','keypress','keyup'].forEach(function(t){
                    try { el.dispatchEvent(new KeyboardEvent(t, opts)); } catch(e) {}
                });
                return true;
            """, inp_el)
            return True
        except Exception as e:
            logger.debug(f"dispatch Enter 실패: {e}")
            return False

    def _submit_search_form(self, inp_el):
        try:
            ok = self.driver.execute_script(
                "var f = arguments[0].closest('form');"
                "if (f) { f.submit(); return true; } return false;",
                inp_el)
            return bool(ok)
        except Exception as e:
            logger.debug(f"form.submit 실패: {e}")
            return False

    def _click_search_icon_near(self, inp_el):
        """검색창 우측의 돋보기(검색) 아이콘을 찾아 클릭.

        dispatchEvent(MouseEvent) 방식으로 WebSquare 이벤트 핸들러를 확실히 자극.
        입력창의 형제 요소부터 탐색 → 부모 DOM 트리 상향 탐색 순으로 진행.
        """
        try:
            clicked = self.driver.execute_script(r"""
                var inp = arguments[0];
                if (!inp) return false;

                function doClick(el) {
                    try {
                        var rect = el.getBoundingClientRect();
                        var cx = rect.left + rect.width  / 2;
                        var cy = rect.top  + rect.height / 2;
                        var opts = {bubbles:true, cancelable:true, view:window,
                                    clientX:cx, clientY:cy};
                        el.dispatchEvent(new MouseEvent('mouseover', opts));
                        el.dispatchEvent(new MouseEvent('mousedown', opts));
                        el.dispatchEvent(new MouseEvent('mouseup',   opts));
                        el.dispatchEvent(new MouseEvent('click',     opts));
                        try { el.click(); } catch(e) {}
                        return true;
                    } catch(e) {
                        try { el.click(); return true; } catch(e2) { return false; }
                    }
                }

                function isVisible(el) {
                    if (!el) return false;
                    var s = window.getComputedStyle(el);
                    if (s.display === 'none' || s.visibility === 'hidden') return false;
                    var r = el.getBoundingClientRect();
                    return (r.width > 0 && r.height > 0);
                }

                function isSearchBtn(el) {
                    if (!el || el === inp) return false;
                    var tn = (el.tagName || '').toLowerCase();
                    if (tn === 'input' || tn === 'select' || tn === 'textarea' ||
                        tn === 'script' || tn === 'style') return false;
                    if (el.contains && el.contains(inp)) return false;
                    if (!isVisible(el)) return false;
                    var t    = ((el.textContent || el.innerText || '').trim());
                    var aria = (el.getAttribute('aria-label') || '');
                    var ttl  = (el.getAttribute('title') || '');
                    var alt  = (el.getAttribute('alt') || '');
                    var cls  = (el.className ? el.className.toString() : '').toLowerCase();
                    var id   = (el.id || '').toLowerCase();
                    // 명시적 텍스트/라벨
                    if (t === '검색' || t === '조회') return true;
                    if (aria.indexOf('검색') !== -1 || aria.toLowerCase().indexOf('search') !== -1) return true;
                    if (ttl.indexOf('검색') !== -1 || alt.indexOf('검색') !== -1) return true;
                    // 클래스/ID 패턴
                    if (cls.indexOf('btn_search') !== -1 || cls.indexOf('btn-search') !== -1) return true;
                    if (cls.indexOf('btn_srch')   !== -1 || cls.indexOf('btn-srch')   !== -1) return true;
                    if (id.indexOf('btn_search')  !== -1 || id.indexOf('btn_srch')    !== -1) return true;
                    // 내부 img alt에 검색
                    var imgs = el.querySelectorAll ? el.querySelectorAll('img') : [];
                    for (var i = 0; i < imgs.length; i++) {
                        if ((imgs[i].getAttribute('alt') || '').indexOf('검색') !== -1) return true;
                    }
                    return false;
                }

                // 1) 입력창 바로 옆 형제 요소 탐색 (가장 빠른 경로)
                var sib = inp.nextElementSibling;
                while (sib) {
                    if (isSearchBtn(sib)) { if (doClick(sib)) return true; }
                    var inner = sib.querySelectorAll ?
                        sib.querySelectorAll('button, a, span, i, em, label') : [];
                    for (var ii = 0; ii < inner.length; ii++) {
                        if (isSearchBtn(inner[ii])) { if (doClick(inner[ii])) return true; }
                    }
                    sib = sib.nextElementSibling;
                }

                // 2) 부모 DOM 트리를 올라가며 탐색 (최대 8단계)
                var container = inp;
                for (var lvl = 0; lvl < 8; lvl++) {
                    container = container.parentElement;
                    if (!container) break;

                    // 명시적 버튼/링크/아이콘 요소
                    var cands = container.querySelectorAll(
                        'button, a, span[role="button"], i, em, label');
                    for (var j = 0; j < cands.length; j++) {
                        if (isSearchBtn(cands[j])) { if (doClick(cands[j])) return true; }
                    }
                    // class/id 패턴 포함 요소 (input 제외)
                    var pats = container.querySelectorAll(
                        '[class*="search"]:not(input), [id*="search"]:not(input),' +
                        '[class*="srch"]:not(input), [id*="srch"]:not(input)');
                    for (var k = 0; k < pats.length; k++) {
                        if (isSearchBtn(pats[k])) { if (doClick(pats[k])) return true; }
                    }
                    // img alt 탐색
                    var imgs2 = container.querySelectorAll('img');
                    for (var m = 0; m < imgs2.length; m++) {
                        if ((imgs2[m].getAttribute('alt') || '').indexOf('검색') !== -1) {
                            var p = imgs2[m].parentElement;
                            if (p && isSearchBtn(p)) { if (doClick(p)) return true; }
                        }
                    }
                }
                return false;
            """, inp_el)
            if clicked:
                logger.info("검색 트리거: 돋보기 아이콘 클릭 (dispatchEvent)")
                return True
        except Exception as e:
            logger.debug(f"검색 아이콘 클릭 중 예외: {e}")
        return False

    # ── 검색 결과 페이지 처리 ─────────────────────────────

    # 결과 페이지 진입을 감지하는 키워드 (사진의 페이지 제목 기준)
    _RESULT_PAGE_KEYWORDS = (
        "부동산 등기사항증명서 열람",
        "부동산 등기사항증명서 발급",
        "간편검색",
        "부동산 소재지번 검색 결과",
        "검색결과",
        "입력하신 검색조건에 대한 결과가 없습니다",
    )

    # "결과 없음" 페이지를 명시적으로 식별하는 키워드.
    _NO_RESULT_KEYWORDS = (
        "입력하신 검색조건에 대한 결과가 없습니다",
        "검색결과가 없습니다",
        "검색 결과가 없습니다",
    )

    def _is_result_page_loaded(self):
        try:
            body = self._get_body_text()
            for kw in self._RESULT_PAGE_KEYWORDS:
                if kw in body:
                    return True
            return False
        except Exception:
            return False

    def _extract_total_count(self):
        """결과 페이지의 '전체 N 건' 카운트를 추출. 못 찾으면 None.

        '결과 없음' 안내 문구가 본문에 명시되면 우선적으로 0 으로 판정.
        """
        try:
            body = self._get_body_text()
            for kw in self._NO_RESULT_KEYWORDS:
                if kw in body:
                    return 0
            # '전체 1 건' 또는 '전체 1건' 모두 매칭. 또한 '전체 0건' '전체 1 건(0건 선택)' 등.
            m = re.search(r"전체\s*([\d,]+)\s*건", body)
            if m:
                num = int(m.group(1).replace(",", ""))
                return num
            # 보조: '검색결과 1건' 같은 페이지
            m2 = re.search(r"검색결과\s*([\d,]+)\s*건", body)
            if m2:
                return int(m2.group(1).replace(",", ""))
        except Exception as e:
            logger.debug(f"카운트 추출 예외: {e}")
        return None

    def _wait_for_search_results(self, timeout=SEARCH_RESULT_WAIT_TIMEOUT,
                                  poll=SEARCH_RESULT_POLL):
        """검색 결과 페이지가 렌더되고 카운트가 산정될 때까지 폴링.

        Returns:
          (loaded: bool, count: Optional[int])
        """
        # 먼저 processbar 가 사라질 때까지 대기 (서버 응답)
        self._wait_until_processbar_gone(max_wait=timeout)
        deadline = time.time() + timeout
        last_count = None
        while time.time() < deadline:
            # 페이지 alert 우선 처리 (검색어 형식 오류 등)
            alert_text = self._handle_alert(accept=True, timeout=0)
            if alert_text:
                logger.warning(f"검색 중 alert: {alert_text[:80]}")
                # alert 가 떴다는 것은 결과 페이지로 진입 못 했다는 신호
                return False, None
            if self._is_result_page_loaded():
                cnt = self._extract_total_count()
                if cnt is not None:
                    last_count = cnt
                    # 안정화 짧게 한 번 더 확인
                    time.sleep(RESULT_STABILIZE_WAIT)
                    cnt2 = self._extract_total_count()
                    if cnt2 is not None:
                        return True, cnt2
                    return True, cnt
            time.sleep(poll)
        # 페이지는 보였으나 카운트를 못 잡은 경우
        if self._is_result_page_loaded():
            return True, last_count
        return False, last_count

    # 결과 표에서 부동산고유번호/구분/표시 추출하는 JS.
    # 헤더 텍스트로 컬럼 인덱스를 찾고, 데이터 행 개수도 함께 반환.
    _EXTRACT_RESULT_JS = r"""
        function visible(el){
            if (!el) return false;
            var s = window.getComputedStyle(el);
            if (s.display === 'none' || s.visibility === 'hidden') return false;
            if (el.offsetParent === null) return false;
            return true;
        }
        function textOf(el){
            return ((el.innerText || el.textContent) || '').replace(/\s+/g, ' ').trim();
        }
        // 후보 테이블: 헤더에 '부동산고유번호' 와 '부동산표시' 가 모두 포함된 것
        var tables = document.querySelectorAll('table');
        var best = null;
        for (var t = 0; t < tables.length; t++) {
            var tb = tables[t];
            if (!visible(tb)) continue;
            var txt = textOf(tb);
            if (txt.indexOf('부동산고유번호') === -1) continue;
            if (txt.indexOf('부동산표시') === -1) continue;
            best = tb;
            break;
        }
        if (!best) return { ok: false, reason: 'table_not_found' };
        // 헤더 행 찾기 (thead 또는 첫 tr)
        var headerCells = best.querySelectorAll('thead th, thead td');
        if (headerCells.length === 0) {
            var firstRow = best.querySelector('tr');
            if (firstRow) headerCells = firstRow.querySelectorAll('th, td');
        }
        if (!headerCells || headerCells.length === 0) {
            return { ok: false, reason: 'no_header' };
        }
        var idxUnique = -1, idxKind = -1, idxRepr = -1;
        for (var h = 0; h < headerCells.length; h++) {
            var ht = textOf(headerCells[h]);
            if (idxUnique === -1 && ht.indexOf('부동산고유번호') !== -1) idxUnique = h;
            else if (idxUnique === -1 && ht.indexOf('고유번호') !== -1) idxUnique = h;
            if (idxKind === -1 && ht.indexOf('부동산구분') !== -1) idxKind = h;
            else if (idxKind === -1 && ht === '구분') idxKind = h;
            if (idxRepr === -1 && ht.indexOf('부동산표시') !== -1) idxRepr = h;
            else if (idxRepr === -1 && ht === '표시') idxRepr = h;
        }
        if (idxUnique === -1 || idxKind === -1 || idxRepr === -1) {
            return { ok: false, reason: 'header_columns_missing',
                     idx: { unique: idxUnique, kind: idxKind, repr: idxRepr } };
        }
        // 데이터 행 수집 (tbody 우선, 없으면 thead 제외한 모든 tr)
        var rows = best.querySelectorAll('tbody tr');
        if (rows.length === 0) {
            rows = best.querySelectorAll('tr');
            // 첫 행이 헤더라면 스킵
            if (rows.length > 0) {
                rows = Array.prototype.slice.call(rows, 1);
            }
        }
        // 보이는 행, 셀 개수가 헤더보다 같거나 더 많은 행만 데이터로 인정.
        var dataRows = [];
        for (var r = 0; r < rows.length; r++) {
            var row = rows[r];
            if (!visible(row)) continue;
            var cells = row.querySelectorAll('td');
            if (cells.length < headerCells.length) continue;
            dataRows.push(cells);
        }
        if (dataRows.length === 0) {
            return { ok: true, count: 0, rows: [] };
        }
        var out = [];
        for (var d = 0; d < dataRows.length; d++) {
            var cells = dataRows[d];
            out.push({
                unique: textOf(cells[idxUnique]),
                kind:   textOf(cells[idxKind]),
                repr:   textOf(cells[idxRepr])
            });
        }
        return { ok: true, count: out.length, rows: out };
    """

    def extract_search_results(self):
        """결과 테이블에서 행 목록을 추출.

        Returns:
          dict: {ok: bool, count: int, rows: [{unique, kind, repr}], reason?: str}
        """
        try:
            result = self.driver.execute_script(
                "return (function(){" + self._EXTRACT_RESULT_JS + "})();"
            )
            if result is None:
                return {"ok": False, "count": 0, "rows": [], "reason": "null_result"}
            return result
        except Exception as e:
            logger.warning(f"결과 추출 JS 실행 실패: {e}")
            return {"ok": False, "count": 0, "rows": [], "reason": f"js_error:{e}"}

    # ── 단건 처리 ────────────────────────────────────────

    def process_single(self, address):
        """주소 한 건에 대해 검색 → 결과 추출. 최대 MAX_RETRY_PER_ITEM 회 재시도.

        Returns:
          dict: {
            status: 'single' | 'multi' | 'none' | 'parse_error' | 'search_error',
            query: str,
            count: int (검색 결과 건수, parse_error 면 None),
            unique, kind, repr: 단건일 때만 값,
            note: 엑셀 비고 컬럼에 기록할 짧은 문자열,
          }
        """
        query, dong_ho_found = parse_address_to_query(address)
        if query is None:
            return {
                "status": "parse_error",
                "query": "",
                "count": None,
                "note": "주소 파싱 실패",
            }

        last_error_note = ""
        for attempt in range(1, MAX_RETRY_PER_ITEM + 1):
            if _pause is not None:
                _pause.checkpoint()

            # 1) 메인으로 이동 (검색은 항상 메인에서)
            if attempt > 1:
                # 재시도 시에는 결과 페이지에 어중간하게 머문 경우를 대비해
                # URL 강제 이동까지 시도 (한 번 실패해도 무시).
                try:
                    if not self._is_truly_on_main():
                        try:
                            self.driver.get(IROS_INDEX_URL)
                        except Exception:
                            pass
                        self._wait_for_page_content()
                    self.go_to_main()
                    time.sleep(0.5)
                except Exception as e:
                    logger.warning(f"메인 이동 중 예외 (무시): {e}")

            inp = self._wait_main_search_box(timeout=8)
            if inp is None:
                try:
                    self.go_to_main()
                    inp = self._wait_main_search_box(timeout=10)
                except Exception:
                    pass

            if inp is None:
                logger.warning(f"  메인 검색창 미감지 (시도 {attempt}/{MAX_RETRY_PER_ITEM})")
                last_error_note = "검색창 미감지"
                time.sleep(RETRY_DELAY_SEC)
                continue

            # 2) 검색 트리거
            ok = self.search_address_on_main(query)
            if not ok:
                logger.warning(f"  검색 트리거 실패 (시도 {attempt}/{MAX_RETRY_PER_ITEM})")
                last_error_note = "검색 트리거 실패"
                time.sleep(RETRY_DELAY_SEC)
                continue

            # 3) 결과 페이지 대기
            loaded, count = self._wait_for_search_results()
            if not loaded:
                logger.warning(
                    f"  결과 페이지 진입 실패 또는 alert 발생 "
                    f"(시도 {attempt}/{MAX_RETRY_PER_ITEM})"
                )
                last_error_note = "결과 페이지 미진입"
                time.sleep(RETRY_DELAY_SEC)
                continue

            # 4) 카운트 기반 분기
            if count == 0:
                logger.info(f"  ✗ 결과 없음 (전체 0 건)")
                return {
                    "status": "none",
                    "query": query,
                    "count": 0,
                    "note": "결과 없음",
                }
            if count is not None and count >= 2:
                logger.info(f"  ⚠ 결과 다건 (전체 {count} 건) → 스킵")
                return {
                    "status": "multi",
                    "query": query,
                    "count": count,
                    "note": f"결과 {count}건 (다건)",
                }

            # count 가 None 이거나 1인 경우 — 테이블 추출로 최종 판정
            extracted = self.extract_search_results()
            if not extracted.get("ok"):
                logger.warning(
                    f"  결과 테이블 추출 실패 ({extracted.get('reason')}) "
                    f"(시도 {attempt}/{MAX_RETRY_PER_ITEM})"
                )
                last_error_note = f"테이블 추출 실패({extracted.get('reason')})"
                time.sleep(RETRY_DELAY_SEC)
                continue

            actual_count = int(extracted.get("count", 0))
            rows = extracted.get("rows", [])
            # 카운트가 None 이었다면 실제 행 수로 결정
            final_count = count if count is not None else actual_count

            if actual_count == 0 or final_count == 0:
                logger.info(f"  ✗ 결과 없음 (테이블 0 행)")
                return {
                    "status": "none",
                    "query": query,
                    "count": 0,
                    "note": "결과 없음",
                }
            if actual_count >= 2 or final_count >= 2:
                logger.info(f"  ⚠ 결과 다건 ({actual_count} 행) → 스킵")
                return {
                    "status": "multi",
                    "query": query,
                    "count": actual_count,
                    "note": f"결과 {actual_count}건 (다건)",
                }

            # 단건 — 추출 성공
            row = rows[0]
            unique = row.get("unique", "").strip()
            kind   = row.get("kind", "").strip()
            repr_  = row.get("repr", "").strip()
            # 너무 공백/줄바꿈이 많을 수 있으므로 정규화
            repr_ = re.sub(r"\s+", " ", repr_)
            logger.info(
                f"  ✓ 단건: {unique} | {kind} | {repr_[:60]}"
                + ("..." if len(repr_) > 60 else "")
            )
            return {
                "status": "single",
                "query": query,
                "count": 1,
                "unique": unique,
                "kind":   kind,
                "repr":   repr_,
                "note":   "단건 추출",
            }

        # 모든 라운드 실패
        return {
            "status": "search_error",
            "query": query,
            "count": None,
            "note": last_error_note or "원인 미상 검색 실패",
        }

    def close(self):
        try:
            if self.driver:
                self.driver.quit()
                logger.info("브라우저 종료")
        except Exception:
            pass
        self.driver = None


# ═══════════════════════════════════════════════════════════
# main
# ═══════════════════════════════════════════════════════════
def main():
    args = parse_args()

    global _overlay, _pause, _popup_watcher
    _pause = PauseController()
    try:
        _overlay = ProgressOverlay.start(controller=_pause)
        _overlay.attach_to_logger(logger)
        _overlay.set_status("초기화 중…")
        _pause.set_state_callback(lambda paused: _overlay.set_pause_indicator(paused))
    except Exception as e:
        _overlay = None
        logger.debug(f"오버레이 초기화 실패 (무시): {e}")

    _popup_watcher = PopupWatcher()
    try:
        _popup_watcher.start()
    except Exception as e:
        logger.debug(f"팝업 워치독 시작 실패 (무시): {e}")

    logger.info("=" * 60)
    logger.info("인터넷등기소 부동산 고유번호 취합 프로그램 v1.5")
    logger.info("=" * 60)

    prevent_sleep()

    if not os.path.exists(args.file):
        logger.error(f"엑셀 파일을 찾을 수 없습니다: {args.file}")
        logger.error(
            f"실행 폴더 ({os.getcwd()}) 안에 '{DEFAULT_FILE}' 파일을 두거나, "
            f"--file <경로> 로 직접 지정해 주세요."
        )
        sys.exit(1)

    try:
        items, cols = load_address_rows(args.file)
    except Exception as e:
        logger.error(f"엑셀 로드 실패: {e}")
        sys.exit(1)

    if not items:
        logger.error("처리할 주소가 엑셀에 없습니다.")
        sys.exit(1)

    total = len(items)
    start_offset = max(args.start_from, 1)
    logger.info(f"총 {total}건 로드 (시작: {start_offset}번째 행부터)")
    logger.info(
        "엑셀 컬럼 매핑: "
        f"주소={get_column_letter(cols['address'])}, "
        f"작업번호={get_column_letter(cols['task_no']) if cols.get('task_no') else '-'}, "
        f"고유번호={get_column_letter(cols['unique'])}, "
        f"구분={get_column_letter(cols['kind'])}, "
        f"표시={get_column_letter(cols['repr'])}, "
        f"비고={get_column_letter(cols['note'])}"
    )

    if start_offset > 1:
        items = items[start_offset - 1:]

    auto = PropertyIdCollector()
    success = 0      # 단건 추출 성공
    none_cnt = 0     # 결과 없음
    multi_cnt = 0    # 다건
    error_cnt = 0    # 파싱/검색 오류

    overlay_status("Phase 1: 검색 시작")
    overlay_progress(0, len(items))

    try:
        auto.start()
        # start() 직후엔 이미 index.jsp(메인) 가 로드된 상태이므로
        # 별도 메인 이동을 호출하지 않는다. 각 검색이 끝난 뒤에만 복귀.

        for idx, item in enumerate(items):
            if _pause is not None:
                _pause.checkpoint()

            row = item["row"]
            task_no = item["task_no"]
            address = item["address"]
            seq_label = task_no or f"행{row}"

            logger.info(f"\n[{idx+1}/{len(items)}] (행{row}, 작업번호={task_no or '-'})")
            logger.info(f"  주소: {address}")
            overlay_status(f"[{idx+1}/{len(items)}] {seq_label}")
            overlay_progress(idx, len(items))

            try:
                result = auto.process_single(address)
            except StopRequested:
                raise
            except Exception as e:
                logger.error(f"  ! 예기치 못한 오류 → 다음 건으로 진행: {e}")
                result = {
                    "status": "search_error",
                    "query": "",
                    "count": None,
                    "note": f"예외:{type(e).__name__}",
                }
                # 메인 페이지로 복구 시도
                try:
                    auto.go_to_main()
                except Exception:
                    pass

            # 엑셀에 결과 기록
            try:
                if result["status"] == "single":
                    write_result_row(
                        args.file, cols, row,
                        unique=result["unique"],
                        kind=result["kind"],
                        repr_=result["repr"],
                        note=result.get("note", "단건 추출"),
                    )
                    success += 1
                elif result["status"] == "multi":
                    write_result_row(args.file, cols, row, note=result["note"])
                    multi_cnt += 1
                elif result["status"] == "none":
                    write_result_row(args.file, cols, row, note=result["note"])
                    none_cnt += 1
                else:
                    write_result_row(args.file, cols, row, note=result["note"])
                    error_cnt += 1
            except Exception as e:
                logger.error(f"  엑셀 기록 실패: {e}")
                error_cnt += 1

            # 다음 건 진입 전 메인 복귀 (마지막 건 제외)
            is_last = (idx == len(items) - 1)
            if not is_last:
                try:
                    auto.go_to_main()
                except Exception as e:
                    logger.warning(f"메인 복귀 중 예외 (무시, 다음 건으로 진행): {e}")
                time.sleep(0.5)

            # 응용 프로그램 오류 휴식 회분 소비
            try:
                if _popup_watcher is not None:
                    _popup_watcher.consume_rest_if_due()
            except StopRequested:
                raise
            except Exception as e:
                logger.debug(f"휴식 회분 소비 중 예외 (무시): {e}")

        logger.info("\n" + "=" * 60)
        logger.info(
            f"[전체 완료] 총 {total}건 | "
            f"단건성공 {success} | 결과없음 {none_cnt} | 다건 {multi_cnt} | 오류 {error_cnt}"
        )
        logger.info("=" * 60)
        overlay_progress(len(items), len(items))
        overlay_status(
            f"✓ 완료 | 성공 {success} / 없음 {none_cnt} / 다건 {multi_cnt} / 오류 {error_cnt}"
        )

    except StopRequested:
        logger.warning("\n" + "=" * 60)
        logger.warning("★ 사용자 요청으로 작업이 중단되었습니다 (오버레이 '중단' 버튼)")
        logger.warning("=" * 60)
        overlay_status("⏹ 중단됨")
        try:
            if _overlay is not None:
                _overlay.set_stopped()
        except Exception:
            pass
    except KeyboardInterrupt:
        logger.warning("\n사용자 중단")
    except Exception as e:
        logger.error(f"오류: {e}", exc_info=True)
    finally:
        try:
            if _popup_watcher is not None and _popup_watcher.dismiss_count > 0:
                logger.warning(
                    f"[팝업워치독] 요약: 자동 닫음 {_popup_watcher.dismiss_count}회, "
                    f"휴식 부여 {_popup_watcher.rest_granted_total}회 / "
                    f"소비 {_popup_watcher.rest_consumed_total}회"
                )
            if _popup_watcher is not None:
                _popup_watcher.stop()
        except Exception:
            pass
        auto.close()
        allow_sleep()
        try:
            if _pause is not None and _pause.is_stopping:
                overlay_status("⏹ 중단됨 — 엔터로 닫기")
            else:
                overlay_status("작업 완료 — 엔터로 닫기")
        except Exception:
            pass


def parse_args():
    p = argparse.ArgumentParser(
        description="인터넷등기소 부동산 고유번호 취합 프로그램 v1.5"
    )
    p.add_argument(
        "--file", "-f", default=DEFAULT_FILE,
        help=f"주소 목록 엑셀 파일 (기본: {DEFAULT_FILE})",
    )
    p.add_argument(
        "--start-from", type=int, default=0,
        help=(
            "N번째 데이터 행부터 처리 (1-based, 헤더 제외 기준). "
            "예) --start-from 5 → 데이터 5번째 줄부터. 미지정/0/1 은 처음부터."
        ),
    )
    return p.parse_args()


if __name__ == "__main__":
    # .py 더블클릭 실행 대응 — 작업 디렉토리를 스크립트 위치로 이동.
    try:
        if getattr(sys, "frozen", False):
            exe_dir = os.path.dirname(os.path.abspath(sys.executable))
        else:
            exe_dir = os.path.dirname(os.path.abspath(__file__))
        if exe_dir:
            os.chdir(exe_dir)
    except Exception:
        pass

    exit_code = 0
    try:
        main()
    except SystemExit as e:
        exit_code = e.code if isinstance(e.code, int) else 1
    except KeyboardInterrupt:
        logger.warning("\n사용자 중단")
        exit_code = 130
    except Exception as e:
        logger.error(f"예기치 못한 오류: {e}", exc_info=True)
        exit_code = 1
    finally:
        try:
            allow_sleep()
        except Exception:
            pass
        try:
            if _popup_watcher is not None:
                _popup_watcher.stop()
        except Exception:
            pass
        try:
            input("\n종료하려면 Enter 를 누르세요...")
        except Exception:
            pass
        try:
            if _overlay is not None:
                _overlay.stop()
        except Exception:
            pass
    sys.exit(exit_code)
