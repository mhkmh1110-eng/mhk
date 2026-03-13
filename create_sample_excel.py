#!/usr/bin/env python3
"""샘플 엑셀 파일 생성 스크립트

프로그램 테스트용 샘플 엑셀 파일을 생성합니다.
"""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


def create_sample():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "부동산고유번호"

    # 스타일 정의
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # 헤더
    headers = ["결과", "부동산고유번호", "소재지(참고)", "비고"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # 샘플 데이터 (실제 고유번호가 아닌 예시)
    sample_data = [
        ("", "1101-2024-012345", "서울시 강남구 역삼동 123-4", ""),
        ("", "1102-2024-067890", "서울시 서초구 서초동 456-7", ""),
        ("", "1301-2024-098765", "경기도 성남시 분당구 정자동 100", ""),
        ("", "2601-2024-054321", "부산시 해운대구 우동 200", ""),
        ("", "2301-2024-011111", "대구시 수성구 범어동 300", ""),
    ]

    for row_num, data in enumerate(sample_data, 2):
        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row_num, column=col, value=value)
            cell.alignment = center_align
            cell.border = thin_border

    # 열 너비 조정
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 15

    wb.save("property_list.xlsx")
    print("샘플 엑셀 파일 생성 완료: property_list.xlsx")


if __name__ == "__main__":
    create_sample()
